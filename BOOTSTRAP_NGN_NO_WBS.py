#!/usr/bin/env python3
r'''Bootstrap NGN's WBS-less General Document as a truthful profile dialect.

Place in cleaner-functions root. Windows CMD:
  py .\BOOTSTRAP_NGN_NO_WBS.py --check
  py .\BOOTSTRAP_NGN_NO_WBS.py --apply --status-date 1405/07/22

The source and working copy remain byte-for-byte identical. No official WBS
is invented; an identity register provides traceable NGN-local document keys.
Existing Cleaner tests, pipeline stages, exceptions and review decisions stay
in place. Stops for human review when the structure is ambiguous.
'''
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parent
PROJECT = 'NGN'
SRC = Path('input/GENERAL_DOCUMENT/NGN/retrieval_20261017/NGN/General - Negin.xlsx')
RECEIPT = Path('input/GENERAL_DOCUMENT/NGN/retrieval_20261017/retrieval_manifest.json')
STAGED = Path('input/GENERAL_DOCUMENT/NGN/NGN_LATEST_GENERAL_DOC_UNAPPROVED.xlsx')
WORKING = Path('working/GENERAL_DOCUMENT/NGN/NGN_GENERAL_DOC_CLEANER_WORKING.xlsx')
PROFILE = Path('profiles/NGN.profile.json')
STAGING_RECEIPT = Path('input/GENERAL_DOCUMENT/NGN/staging_receipt.json')
IDENTITIES = Path('input/GENERAL_DOCUMENT/NGN/ngn_document_identity_register.json')
IDENTIFIER_TEST = Path('tests/identifier_check.py')
TEST_FILE = Path('unit_tests/test_ngn_optional_wbs.py')
ANCHOR = '        if key == "wbs" or strategy == "NEAREST_VALID_ABOVE_APPEND_SEQUENCE":'
OPTIONAL = '        if key == "wbs" and not fdef.get("required", True):\n            continue\n'
TEST_SOURCE = '''from tests.identifier_check import _sequence_fields


def test_ngn_explicit_optional_wbs_is_not_synthesized():
    fields = {"wbs": {"required": False, "cleaning": {"strategy": "NEAREST_VALID_ABOVE_APPEND_SEQUENCE"}}}
    assert _sequence_fields(fields) == {}


def test_required_wbs_keeps_existing_family_repair():
    fields = {"wbs": {"required": True, "cleaning": {"strategy": "NEAREST_VALID_ABOVE_APPEND_SEQUENCE"}}}
    assert "wbs" in _sequence_fields(fields)


def test_other_optional_sequence_fields_keep_their_explicit_policy():
    fields = {"code": {"required": False, "cleaning": {"strategy": "NEAREST_VALID_ABOVE_APPEND_SEQUENCE"}}}
    assert "code" in _sequence_fields(fields)
'''


def digest(p: Path) -> str:
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda: f.read(1<<20),b''):
            h.update(b)
    return h.hexdigest()


def read_json(p: Path) -> dict:
    return json.loads(p.read_text(encoding='utf-8'))


def render_json(obj: dict) -> str:
    return json.dumps(obj, indent=2, ensure_ascii=False, default=str)+'\n'


def check_current(target: Path, content: str):
    if target.exists() and target.read_text(encoding='utf-8') != content:
        raise RuntimeError(f'Existing file differs, refusing to overwrite: {target}')


def copy_verified(src: Path, dest: Path, sha: str) -> str:
    if dest.exists():
        if digest(dest) != sha:
            raise RuntimeError(f'Existing file differs, refusing overwrite: {dest}')
        return 'EXISTING_VERIFIED'
    dest.parent.mkdir(parents=True, exist_ok=True)
    fd, tempname = tempfile.mkstemp(prefix='.ngn_',dir=dest.parent)
    os.close(fd)
    temp=Path(tempname)
    try:
        shutil.copy2(src,temp)
        if digest(temp)!=sha:
            raise RuntimeError('SHA256 changed while staging')
        temp.rename(dest)
    finally:
        temp.unlink(missing_ok=True)
    return 'CREATED'


def source_and_sha():
    source=ROOT/SRC
    manifest=ROOT/RECEIPT
    if not source.is_file() or not manifest.is_file():
        raise RuntimeError('Verified NGN retrieval not found; expected '+str(source))
    rows=[x for x in read_json(manifest).get('projects',[]) if x.get('project')=='NGN']
    if len(rows)!=1:
        raise RuntimeError('NGN retrieval manifest must contain exactly one selected workbook')
    selection=rows[0].get('selected') or {}
    expected=str(selection.get('sha256') or '').lower()
    if not re.fullmatch(r'[0-9a-f]{64}',expected) or digest(source)!=expected:
        raise RuntimeError('NGN retrieved source SHA mismatch; no files changed')
    if Path(str(selection.get('name') or '')).name != source.name:
        raise RuntimeError('NGN retrieval manifest and selected workbook name disagree')
    return source, expected, selection


def rows_and_provenance(source: Path, sha: str):
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    w=load_workbook(source,read_only=False,data_only=False)
    try:
        if w.sheetnames!=['MOBILE SS']:
            raise RuntimeError(f'NGN sheet selection needs human review: {w.sheetnames!r}')
        ws=w['MOBILE SS']
        if ws.max_column < 6:
            raise RuntimeError('Expected D/E/F transaction board columns, none found')
        header=' '.join(str(ws.cell(1,c).value or '') for c in (4,5,6)).lower()
        if not ('send' in header and 'receive' in header):
            raise RuntimeError('D/E correspondence header not confirmed; stop for profile review')
        if any(not isinstance(ws.cell(r,3),MergedCell) and str(ws.cell(r,3).value or '').strip() for r in range(3,ws.max_row+1)):
            raise RuntimeError('Column C is not empty: NGN may have a document code or WBS; review before staging')
        docs=[]
        counts=Counter()
        for r in range(3,ws.max_row+1):
            cell=ws.cell(r,2)
            if isinstance(cell,MergedCell):
                continue
            name=str(cell.value or '').strip()
            if not name:
                continue
            # B ranges covered by a full-width separator anchored in A are MergedCell; skipped above.
            normalized=' '.join(name.casefold().split())
            counts[normalized]+=1
            occurrence=counts[normalized]
            stable_seed=f'MOBILE SS|{normalized}|{occurrence}'
            key='NGN-DOC-'+hashlib.sha256(stable_seed.encode('utf8')).hexdigest()[:16].upper()
            docs.append({'internal_document_key':key,'key_type':'GENERATED_NON_WBS',
                         'source_sheet':'MOBILE SS','source_name_cell':f'B{r}',
                         'source_wbs_cell':f'C{r}','source_wbs':None,
                         'document_name':name,'original_anchor_row':r,
                         'duplicate_name_ordinal':occurrence,
                         'source_name_sha256':hashlib.sha256(name.encode('utf8')).hexdigest(),
                         'status':'UNVERIFIED_AS_OFFICIAL_IDENTIFIER'})
        if len(docs)<8:
            raise RuntimeError(f'Only {len(docs)} named document anchors; profile review required')
        if len({d['internal_document_key'] for d in docs}) != len(docs):
            raise RuntimeError('Generated key collision; requires human review')
        # Stable local lookup only; no ID is written into the Excel file.
        return {'schema':'cleaner.document_identifier_provenance/0.1',
                'project':'NGN','family':'GENERAL_DOCUMENT','source_sha256':sha,
                'source_sheet':'MOBILE SS','official_wbs_present':False,
                'generated_identifiers_are_official_wbs':False,
                'generation_rule':'SHA256(sheet|normalized_name|occurrence), first 16 hex chars',
                'warning':'Do not publish internal_document_key as official project WBS; keyed by source name and occurrence, pending business ID reconciliation.',
                'duplicate_document_names':sum(n-1 for n in counts.values() if n>1),
                'document_count':len(docs),'documents':docs}
    finally:
        w.close()


def layout_override() -> dict:
    return {'identifier_layout':{'fields':[
        {'key':'document_type','label':'Document Type','column':1,'required':False},
        {'key':'name','label':'Name','column':2,'required':True},
        {'key':'wbs','label':'Source WBS (ABSENT / OPTIONAL)','column':3,
         'required':False,'validation':{'kind':'NONE'},'cleaning':{}},
    ],'boundary_after_col':3,'data_start_col':4,
       'signal_fields':['name'], 'boundary_fields':['name'], 'separator_fields':['name']}}


def build_profile(source: Path):
    from functions.cleaner.bootstrap import build_general_document_profile
    p=ROOT/PROFILE
    _, prof=build_general_document_profile(project_id='NGN',workbook_path=source,
        base_profile_path=ROOT/'profiles/families/GENERAL_DOCUMENT.base.json',
        output_path=p,sheets=['MOBILE SS'],profile_overrides=layout_override(),overwrite=False)
    prof['data_model']['identity_provenance']={
        'source_wbs_present':False,'wbs_required':False,
        'internal_identifier_register':str(ROOT/IDENTITIES),
        'do_not_publish_as_official_wbs':True}
    p.write_text(render_json(prof),encoding='utf8')
    return prof


def main(argv=None)->int:
    ap=argparse.ArgumentParser(description=__doc__)
    op=ap.add_mutually_exclusive_group(required=True)
    op.add_argument('--check',action='store_true')
    op.add_argument('--apply',action='store_true')
    ap.add_argument('--status-date',default='1405/07/22')
    ap.add_argument('--no-run',action='store_true',help='Stage and configure only')
    ap.add_argument('--run-id',default=None,help='Resume NGN existing run')
    a=ap.parse_args(argv)
    if not re.fullmatch(r'14\d{2}/\d{2}/\d{2}',a.status_date):
        ap.error('Invalid Jalali status date')
    os.chdir(ROOT)
    ident=ROOT/IDENTIFIER_TEST
    if not ident.is_file() or not (ROOT/'cleaner.package.json').is_file():
        raise RuntimeError('Place script directly in cleaner-functions repository root')
    code=ident.read_text(encoding='utf8')
    if OPTIONAL+ANCHOR in code:
        new_code=code
        change='ALREADY_APPLIED'
    elif code.count(ANCHOR)==1:
        new_code=code.replace(ANCHOR,OPTIONAL+ANCHOR,1)
        change='PATCH_REQUIRED'
    else:
        raise RuntimeError('Optional-WBS code anchor changed; stop rather than guess a patch')
    source, sha, selection=source_and_sha()
    provenance=rows_and_provenance(source,sha)
    print(f'[SOURCE] verified sha256={sha} | {source.name}')
    print(f'[PROFILE] MOBILE SS / Name=B / WBS=ABSENT(optional C) / Data=D:X')
    print(f'[DOCUMENTS] {provenance["document_count"]} named anchors, duplicate names={provenance["duplicate_document_names"]}')
    print(f'[RULE] Optional-WBS compatibility: {change}')
    for dest in (ROOT/STAGED,ROOT/WORKING):
        if dest.exists() and digest(dest)!=sha:
            raise RuntimeError(f'Existing source or working workbook differs: {dest}')
    profile=ROOT/PROFILE
    if profile.exists():
        existing=read_json(profile)
        fields=existing.get('data_model',{}).get('family_fields',[])
        wbs=next((f for f in fields if f.get('key')=='wbs'),None)
        if existing.get('profile_name')!='NGN' or not wbs or wbs.get('required',True):
            raise RuntimeError(f'NGN profile already exists but is not this optional-WBS profile: {profile}')
    for target, content in [(ROOT/TEST_FILE,TEST_SOURCE),(ROOT/IDENTITIES,render_json(provenance))]:
        check_current(target,content)
    # Ensure receipt is stable/idempotent even on repeated execution.
    receipt={'schema':'cleaner.general_document_staging/0.1','project':'NGN','family':'GENERAL_DOCUMENT',
         'state':'PENDING_SOURCE_AND_PROFILE_REVIEW','source_name':source.name,
         'source_original_path':selection.get('source'), 'source_sha256':sha,
         'staged_source':str(ROOT/STAGED),'working_copy':str(ROOT/WORKING),
         'profile_path':str(ROOT/PROFILE),'identifier_provenance':str(ROOT/IDENTITIES),
         'note':'WBS intentionally absent in NGN source; no fabricated official WBS. Structure still subject to review.'}
    if (ROOT/STAGING_RECEIPT).exists() and read_json(ROOT/STAGING_RECEIPT).get('source_sha256')!=sha:
        raise RuntimeError('Existing staging receipt points to different source')
    if a.check:
        print('[CHECK] Compatible; no files modified. Ready for NGN no-WBS profile and review-driven run.')
        return 0
    # Save existing core file before modification. Existing data and run files are never overwritten.
    if change=='PATCH_REQUIRED':
        backup=ROOT/('_ngn_optional_wbs_backup_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
        backup.mkdir(exist_ok=False)
        shutil.copy2(ident,backup/ident.name)
        ident.write_text(new_code,encoding='utf8')
        print(f'[PATCH] Optional WBS supported, backup={backup}')
    for dest in (ROOT/STAGED, ROOT/WORKING):
        print(f'[STAGE] {dest.name}: {copy_verified(source,dest,sha)}')
    if not profile.exists():
        build_profile(ROOT/STAGED)
        print(f'[PROFILE] created {profile}')
    else:
        print('[PROFILE] existing compatible profile retained')
    for target, content in [(ROOT/TEST_FILE,TEST_SOURCE),(ROOT/IDENTITIES,render_json(provenance))]:
        if not target.exists():
            target.parent.mkdir(parents=True,exist_ok=True)
            target.write_text(content,encoding='utf8')
    receiptpath=ROOT/STAGING_RECEIPT
    if not receiptpath.exists():
        receiptpath.parent.mkdir(parents=True,exist_ok=True)
        receiptpath.write_text(render_json(receipt),encoding='utf8')
    print('[PROVENANCE] Original workbook unchanged; NGN local document-key register created.')
    if a.no_run:
        print('[READY] py -m pytest -q --ignore=payload --import-mode=importlib')
        return 0
    runner=ROOT/'RUN_GENERAL_DOCUMENT_V03.py'
    if not runner.is_file():
        print('[READY] Runner V03 missing. Stage complete; install runner first.')
        return 2
    # Choose existing one unambiguously; never silently fork repeated run state.
    matches=[]
    for p in (ROOT/'runs').glob('ngn_*/run_manifest.json'):
        try:
            o=read_json(p); wb=o.get('workbook') or {}
            if (o.get('profile_name')=='NGN' and
                Path(str(wb.get('path') or '')).resolve()==(ROOT/WORKING).resolve() and
                wb.get('sha256_at_run_start')==sha):
                matches.append(p.parent.name)
        except (OSError,ValueError,KeyError):
            continue
    if a.run_id:
        if a.run_id not in matches:
            raise RuntimeError(f'Run ID {a.run_id} does not match verified NGN source')
        mode=['--run-id',a.run_id]
    elif len(matches)==1:
        mode=['--run-id',matches[0]]
        print(f'[RESUME] {matches[0]}')
    elif len(matches)>1:
        raise RuntimeError(f'Multiple NGN runs match: {matches}; rerun with --run-id')
    else:
        mode=['--new-run']
    print('[START] Review-driven NGN pipeline; no invented WBS.',flush=True)
    env=os.environ.copy()
    for v in ('OPENBLAS_NUM_THREADS','OMP_NUM_THREADS','MKL_NUM_THREADS','NUMEXPR_NUM_THREADS'):
        env[v]='1'
    return subprocess.call([sys.executable,str(runner),'--project','NGN',*mode,
                           '--workbook',str(ROOT/WORKING),'--status-date',a.status_date],cwd=ROOT,env=env)

if __name__=='__main__':
    try:
        sys.exit(main())
    except Exception as exc:
        print('[STOP]',str(exc),file=sys.stderr)
        sys.exit(1)
