#!/usr/bin/env python3
"""Run-local, operator-confirmed BFG receive-status correction (source read-only).

From cleaner-functions root:
  py .\SET_BFG_MISSING_STATUS_NOT_APPROVED.py --check
  py .\SET_BFG_MISSING_STATUS_NOT_APPROVED.py --apply

The 1404/04/29 RECEIVE DATE at K69 is never changed. The blank corresponding
status merged cell gets an audited logical override of N.A (Not Approved).
"""
from __future__ import annotations
import argparse
from copy import deepcopy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
import sys

PROJECT = 'BFG'
RUN_ID = 'bfg_20261017T130419'
SHEET = 'BAAFGH'
DATE_CELL = 'K69'
DATE_VALUE = '1404/04/29'
STATUS_COLUMN = 12  # L, adjacent to K receive column in the 3-column block
STATUS_VALUE = 'N.A'
REVIEW_ID = 'USER_CONFIRMED_BFG_L69_RECEIVE_STATUS_NOT_APPROVED'
SOURCE_SHA = '444b798178f71e127fccd967f4fb6d7a8029b5f7eaddc87275a1a79ce1fdecf7'
ROOT = Path(__file__).resolve().parent


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def guard(test, message):
    if not test:
        raise RuntimeError(message)


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1048576), b''):
            h.update(block)
    return h.hexdigest()


def write_json(path, obj):
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
    tmp.replace(path)


def prepare(root):
    run = root / 'runs' / RUN_ID
    guard(run.is_dir(), f'Missing run: {run}')
    manifest_path, state_path = run / 'run_manifest.json', run / 'resolved_state.json'
    card_path, schema_path = run / 'card_maps.json', run / 'slot_schema.json'
    registry_path = root / 'exceptions' / 'BFG.exceptions.json'
    for p in (manifest_path, state_path, card_path, schema_path):
        guard(p.is_file(), f'Required artifact missing: {p}')
    manifest, state, cards, schema = [read(p) for p in (manifest_path, state_path, card_path, schema_path)]
    guard(manifest.get('profile_name') == PROJECT, 'Wrong project in run manifest')
    guard(state.get('run_id') == RUN_ID and state.get('profile_name') == PROJECT, 'Wrong run resolved state')
    book = Path(str((manifest.get('workbook') or {}).get('path') or ''))
    guard(book.is_file(), f'Working workbook missing: {book}')
    guard(book.name == 'BFG_GENERAL_DOC_CLEANER_WORKING.xlsx', 'Unexpected workbook name')
    guard(digest(book) == SOURCE_SHA, 'Workbook differs from verified staged BFG source; stop and inspect')
    sheet_schema = (schema.get('sheets') or {}).get(SHEET) or {}
    start = int(sheet_schema.get('data_start_col') or 0)
    width = int(sheet_schema.get('block_width') or 0)
    guard(width == 3 and start > 0, f'Unexpected transaction width/start: {width}/{start}')
    guard((11 - start) % width == 1 and (STATUS_COLUMN - start) % width == 2, 'K/L not RECEIVE/STATUS columns under inferred schema')
    roles = sheet_schema.get('roles_by_offset') or {}
    guard(str(roles.get('2') or '').lower() == 'status', 'Inferred third column is not a status')
    sheet_cards = ((cards.get('sheets') or {}).get(SHEET) or {}).get('cards') or []
    matching_cards = [c for c in sheet_cards if int(c.get('start_row') or 0) <= 69 <= int(c.get('end_row') or 0)]
    guard(len(matching_cards) == 1, f'Expected one document card on row 69; found {len(matching_cards)}')
    possible_slots = [s for s in (matching_cards[0].get('slots') or []) if int(s.get('col_start') or 0) == STATUS_COLUMN and int(s.get('col_end') or 0) == STATUS_COLUMN]
    guard(len(possible_slots) == 1, f'Expected one status slot in L column for this card; found {len(possible_slots)}')
    slot = possible_slots[0]
    target = str(slot.get('address') or '').upper()
    guard(target in {'L69', 'L69:L70'}, f'Unexpected status slot: {target}')
    guard(slot.get('value') in (None, ''), f'Mapped status not empty: {slot.get("value")!r}')
    from openpyxl import load_workbook
    wb = load_workbook(book, read_only=False, data_only=False)
    try:
        guard(SHEET in wb.sheetnames, 'BAAFGH sheet not found')
        ws = wb[SHEET]
        guard(str(ws[DATE_CELL].value).strip() == DATE_VALUE, f'Receive date changed: {ws[DATE_CELL].value!r}')
        guard(ws['L69'].value in (None, ''), f'Status slot is populated: {ws["L69"].value!r}')
        if target == 'L69:L70':
            guard(any(str(rng) == target for rng in ws.merged_cells.ranges), 'Expected L69:L70 merged status range missing')
        else:
            guard(not any('L69' in rng for rng in ws.merged_cells.ranges), 'Status is merged but schema mapped individual cell')
    finally:
        wb.close()
    catalog = read(root / 'configs' / 'status_library.json')
    defined = (catalog.get('statuses') or {}).get(STATUS_VALUE) or {}
    guard(defined.get('meaning') == 'Not Approved' and defined.get('responsibility_after_receive') == 'COMPANY', 'N.A must be defined as Not Approved, COMPANY responsibility')
    key = f'{SHEET}!{target}'
    old = (state.get('value_overrides') or {}).get(key)
    if old is not None:
        guard(old.get('source_review_id') == REVIEW_ID and old.get('value') == STATUS_VALUE, f'Conflicting status override: {old!r}')
    guard('BAAFGH!K69' not in (state.get('value_overrides') or {}), 'Unsafe K69 DATE override still present; run RECOVER_BFG_K69.py first')
    reg = read(registry_path) if registry_path.is_file() else None
    cases = [] if reg is None else [c for c in reg.get('cases', []) if c.get('last_review_id', '').startswith('FIELD_VALIDATION_BAAFGH_K69_') or (c.get('source') or {}).get('first_review_id', '').startswith('FIELD_VALIDATION_BAAFGH_K69_')]
    guard(len(cases) <= 1, 'Multiple K69 exception cases; refuse ambiguous closure')
    return locals()


def main():
    p = argparse.ArgumentParser()
    modes = p.add_mutually_exclusive_group(required=True)
    modes.add_argument('--check', action='store_true')
    modes.add_argument('--apply', action='store_true')
    args = p.parse_args()
    try:
        ctx = prepare(ROOT)
        print(f'[CHECK] Receive date: {SHEET}!{DATE_CELL} = {DATE_VALUE} (unchanged)')
        print(f'[CHECK] Status slot: {SHEET}!{ctx["target"]} = blank -> {STATUS_VALUE} (Not Approved)')
        print(f'[CHECK] BFG run: {RUN_ID}; existing case(s): {len(ctx["cases"])}')
        if args.check:
            print('[CHECK] No files modified. Ready to apply run-local correction.')
            return 0
        now = datetime.now().isoformat(timespec='seconds')
        state = ctx['state']
        old = (state.get('value_overrides') or {}).get(ctx['key'])
        if old and old.get('source_review_id') == REVIEW_ID:
            print('[OK] Already applied; no duplicate correction made.')
            return 0
        source = {'test_id': 'OPERATOR_CONFIRMED_RECEIVE_STATUS', 'sheet': SHEET, 'receive_date_source': DATE_CELL, 'receive_date': DATE_VALUE}
        reason = 'Operator confirmed missing receive status is N.A (Not Approved), not a blank/unknown status. Preserve K69 date.'
        decision = {
            'action': 'SET_VALUE',
            'parameters': {'sheet': SHEET, 'target': ctx['target'], 'value': STATUS_VALUE, 'expected_old_value': ''},
            'reason': reason, 'reviewer': 'operator-confirmed-in-chat', 'decided_at': now,
            'execution_kind': 'WORKBOOK_CHANGE', 'execution_state': 'PENDING'
        }
        state.setdefault('value_overrides', {})[ctx['key']] = {
            'sheet': SHEET, 'target': ctx['target'], 'value': STATUS_VALUE, 'action': 'SET_VALUE',
            'expected_old_value': '', 'source_review_id': REVIEW_ID, 'reason': reason, 'decided_at': now,
        }
        state.setdefault('decision_audit', []).append({
            'review_id': REVIEW_ID, 'source': source,
            'location': {'sheet': SHEET, 'target': ctx['target']},
            'decision': decision, 'ingested_at': now,
        })
        state.setdefault('review_corrections', []).append({
            'review_id': REVIEW_ID, 'corrected_at': now,
            'linked_original_finding': 'FIELD_VALIDATION_BAAFGH_K69_273d98c30682',
            'status_cell': ctx['target'], 'receive_date_preserved': DATE_VALUE,
            'notes': 'No family-wide default: individual operator-confirmed BFG status.'
        })
        state['updated_at'] = now
        receipt = {
            'schema': 'cleaner.operator_status_confirmation/0.1', 'run_id': RUN_ID,
            'source_sha256': SOURCE_SHA, 'sheet': SHEET, 'receive_date_source': DATE_CELL,
            'receive_date_value': DATE_VALUE, 'status_cell': ctx['target'],
            'original_status': None, 'confirmed_status': STATUS_VALUE,
            'meaning': 'Not Approved', 'reason': reason,
            'source_workbook_modified': False, 'confirmed_at': now, 'review_id': REVIEW_ID,
        }
        if ctx['cases']:
            case = ctx['cases'][0]
            case['case_state'] = 'RESOLVED'
            case['workflow_status'] = 'OPERATOR_CONFIRMED_STATUS'
            case['classification'] = 'CORRECTED_IN_RESOLVED_STATE'
            case['resolved_at'] = now
            case.setdefault('history', []).append({
                'at': now, 'event': 'OPERATOR_CONFIRMED_STATUS', 'review_id': REVIEW_ID,
                'note': f'{SHEET}!{ctx["target"]} set to {STATUS_VALUE} in run overlay; date preserved.'})
        backup = ROOT / ('_bfg_status_confirm_backup_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
        guard(not backup.exists(), f'Backup exists: {backup}')
        backup.mkdir()
        for path in [ctx['state_path'], ctx['registry_path']]:
            if path.is_file():
                shutil.copy2(path, backup / path.name)
        write_json(ctx['state_path'], state)
        if ctx['reg'] is not None and ctx['cases']:
            write_json(ctx['registry_path'], ctx['reg'])
        write_json(ctx['run'] / 'bfg_k69_operator_status_confirmation.json', receipt)
        print(f'[APPLIED] {SHEET}!{ctx["target"]}: blank -> N.A (Not Approved), run overlay only')
        print(f'[PRESERVED] {SHEET}!{DATE_CELL}: {DATE_VALUE}; workbook untouched')
        print(f'[BACKUP] {backup}')
        print(f'[NEXT] py .\\RUN_GENERAL_DOCUMENT_V03.py --project BFG --run-id "{RUN_ID}" --status-date 1405/07/22')
        return 0
    except Exception as e:
        print(f'[STOP] {e}. Do not publish previous engineering report.')
        return 1

if __name__ == '__main__':
    sys.exit(main())
