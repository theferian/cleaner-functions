#!/usr/bin/env python3
r"""Stage a verified General Document retrieval and launch the reviewed Cleaner runner.

Place in cleaner-functions root, next to RUN_GENERAL_DOCUMENT_V02.py.

Examples (Windows CMD):
  py .\STAGE_AND_RUN_GENERAL_DOC.py --project BFG --retrieval "input\GENERAL_DOCUMENT\BFG\retrieval_20261017" --status-date 1405/07/22
  py .\STAGE_AND_RUN_GENERAL_DOC.py --project BFG --retrieval "input\GENERAL_DOCUMENT\BFG\retrieval_20261017" --inspect-only

Only unapproved, hash-verified candidate copies are staged. Source approvals and business sign-off remain separate.
Never edits network originals. Ambiguous identifier layouts are a review gate, not a guess.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from tempfile import NamedTemporaryFile

ROOT = Path(__file__).resolve().parent
FAMILY = "GENERAL_DOCUMENT"
WBS_RE = re.compile(r"^\d+(?:\.\d+)+$")
NON_DOCUMENT_NAMES = ("transmittal", "weekly", "coverage_audit", "answer to comment")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def json_read(path: Path) -> dict:
    obj = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(obj, dict):
        raise ValueError(f"Not a JSON object: {path}")
    return obj


def safe_copy(src: Path, dst: Path, expected: str) -> str:
    if dst.exists():
        if sha256(dst) != expected:
            raise RuntimeError(f"Different file already exists; will not overwrite {dst}")
        return "EXISTING_MATCHED"
    dst.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(mode="wb", prefix=".stage_", dir=dst.parent, delete=False) as handle:
        tmp = Path(handle.name)
    try:
        shutil.copyfile(src, tmp)
        if sha256(tmp) != expected:
            raise RuntimeError(f"Copy hash mismatch: {dst}")
        if dst.exists():
            if sha256(dst) != expected:
                raise RuntimeError(f"Concurrent change to {dst}")
            return "EXISTING_MATCHED"
        tmp.rename(dst)
        return "CREATED"
    finally:
        tmp.unlink(missing_ok=True)


def load_selected(retrieval: Path, project: str) -> tuple[Path, dict, str]:
    manifest = json_read(retrieval / "retrieval_manifest.json")
    entries = [e for e in manifest.get("projects", []) if str(e.get("project", "")).upper() == project]
    if len(entries) != 1 or entries[0].get("status") != "COPIED_REVIEW_REQUIRED":
        raise RuntimeError("No unique, hash-verified candidate in retrieval manifest")
    selected = entries[0].get("selected") or {}
    name = selected.get("name", "")
    if not name or Path(name).name != name or Path(name).suffix.lower() != ".xlsx":
        raise RuntimeError(f"Candidate is not an .xlsx General Document: {name!r}")
    if any(token in name.casefold() for token in NON_DOCUMENT_NAMES):
        raise RuntimeError(f"Candidate looks like a non-General-Document file: {name}")
    source = retrieval / project / name
    if not source.is_file():
        raise FileNotFoundError(f"Verified retrieval copy not found: {source}")
    expected = str(selected.get("sha256") or "").lower()
    if not re.fullmatch(r"[0-9a-f]{64}", expected) or sha256(source) != expected:
        raise RuntimeError(f"Retrieval SHA-256 mismatch for {source}")
    return source, selected, expected


def inspect_layout(book: Path) -> tuple[list[str], dict | None, list[str]]:
    """Conservative, evidence-based identifier detection; never synthesize identifiers blindly."""
    from openpyxl import load_workbook
    from functions.cleaner.bootstrap import inspect_general_document_workbook

    probes = inspect_general_document_workbook(book)
    selected = [p.sheet for p in probes if p.selected]
    report = [f"{p.sheet}: rows={p.max_row}, cols={p.max_col}, identifier_rows={p.identifier_rows}, board_cells={p.transaction_cells}, selected={p.selected}" for p in probes]
    if not selected or len(selected) > 3:
        return selected, None, report + ["Selected data-sheet count outside safe range (1–3)."]
    wb = load_workbook(book, read_only=False, data_only=False)
    try:
        per_sheet = []
        for sheet in selected:
            ws = wb[sheet]
            maxr = min(ws.max_row, 450)
            scores: list[tuple[int, int]] = []
            for col in range(2, min(ws.max_column, 8) + 1):
                count = sum(bool(WBS_RE.fullmatch(str(ws.cell(row, col).value or "").strip())) for row in range(3, maxr + 1))
                scores.append((count, col))
            scores.sort(reverse=True)
            report.append(f"{sheet}: WBS candidate columns=" + ", ".join(f"{col}:{count}" for count, col in scores[:5]))
            if not scores or scores[0][0] < 8 or (len(scores) > 1 and scores[1][0] > scores[0][0] * .35):
                return selected, None, report + [f"{sheet}: WBS column not sufficiently distinctive; human profile review needed."]
            wbs = scores[0][1]
            if wbs < 3 or wbs > 6:
                return selected, None, report + [f"{sheet}: WBS candidate column {wbs} outside supported range."]
            # The name column is chosen from human-language, descriptive cells to the left
            # of WBS, excluding column A which is normally a short category identifier.
            name_scores = []
            for col in range(2, wbs):
                values = [str(ws.cell(r,col).value or "").strip() for r in range(3,maxr+1)]
                descriptive = [v for v in values if len(v)>=7 and any(ch.isalpha() for ch in v)
                               and not WBS_RE.fullmatch(v)]
                # multi-word or long field name; code columns receive a lower score
                score = sum(min(len(v),70) * (1.5 if " " in v else 1) for v in descriptive)
                name_scores.append((score,col,len(descriptive)))
            name_scores.sort(reverse=True)
            report.append(f"{sheet}: descriptive-name candidate columns=" + ", ".join(f"{c}:{n}/{s:.0f}" for s,c,n in name_scores))
            if not name_scores or name_scores[0][2] < 8 or (len(name_scores)>1 and name_scores[1][0] > .85*name_scores[0][0]):
                return selected, None, report + [f"{sheet}: Name column ambiguous; human profile review needed."]
            name = name_scores[0][1]
            per_sheet.append((sheet,name,wbs))
        if len({(n,w) for _,n,w in per_sheet}) != 1:
            return selected,None,report+[f"Identifier layout differs across selected sheets: {per_sheet}"]
        _, name, wbs = per_sheet[0]
        return selected, {"name_col":name,"wbs_col":wbs}, report
    finally:
        wb.close()


def make_overrides(layout: dict) -> dict:
    name, wbs = int(layout["name_col"]), int(layout["wbs_col"])
    fields = []
    for col in range(1, wbs + 1):
        if col == name:
            key, label, required = "name", "Name", True
        elif col == wbs:
            key, label, required = "wbs", "WBS", True
        elif col == 1:
            key, label, required = "document_type", "Document Type", False
        elif col == wbs - 1 and col > name:
            key, label, required = "code", "Code", False
        else:
            key, label, required = f"identifier_{col}", f"Identifier {col}", False
        fields.append({"key": key, "label": label, "column": col, "required": required})
    return {"identifier_layout": {"fields": fields, "boundary_after_col": wbs, "data_start_col": wbs+1,
                                  "signal_fields":[f["key"] for f in fields],
                                  "boundary_fields":["name","wbs"],
                                  "separator_fields":["name","wbs"]}}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--project",required=True)
    ap.add_argument("--retrieval",required=True,help="Folder containing retrieval_manifest.json and PROJECT/workbook.xlsx")
    ap.add_argument("--status-date",default="1405/07/22",help="Jalali reporting date")
    ap.add_argument("--inspect-only",action="store_true",help="Validate source and inspect layout; do not write or run")
    ap.add_argument("--run-id",default=None,help="Resume the exact existing run; otherwise discover a single matching run")
    args=ap.parse_args(argv)
    project=args.project.strip().upper()
    if not re.fullmatch(r"[A-Z0-9_]+",project):
        ap.error("invalid project")
    if not re.fullmatch(r"14\d{2}/\d{2}/\d{2}",args.status_date):
        ap.error("invalid Jalali date")
    if not (ROOT/"pipeline"/"project_pipeline.py").is_file():
        raise RuntimeError("Place script in cleaner-functions repository root")
    retrieval=Path(args.retrieval)
    if not retrieval.is_absolute():
        retrieval=ROOT/retrieval
    source, selected, expected=load_selected(retrieval,project)
    print(f"[SOURCE] {project} {selected.get('name')} | sha256={expected}")
    sheets, layout, report=inspect_layout(source)
    for line in report:
        print("[INSPECT]",line)
    if layout is None:
        print("[REVIEW REQUIRED] Identifier layout not proven. Nothing staged or run. Provide [INSPECT] output to configure profile.")
        return 2
    print(f"[LAYOUT] sheets={sheets} name={layout['name_col']} WBS={layout['wbs_col']} data_start={layout['wbs_col']+1} (PROPOSED)")
    if args.inspect_only:
        print("[CHECK] Read-only inspection complete; no files changed.")
        return 0
    staged = ROOT/"input"/FAMILY/project/f"{project}_LATEST_GENERAL_DOC_UNAPPROVED.xlsx"
    working = ROOT/"working"/FAMILY/project/f"{project}_GENERAL_DOC_CLEANER_WORKING.xlsx"
    profile_path = ROOT/"profiles"/f"{project}.profile.json"
    receipt_path = ROOT/"input"/FAMILY/project/"staging_receipt.json"
    for dest in [staged,working]:
        if dest.exists() and sha256(dest)!=expected:
            raise RuntimeError(f"Cannot overwrite different existing workbook: {dest}")
    if receipt_path.exists() and json_read(receipt_path).get("source_sha256")!=expected:
        raise RuntimeError(f"Existing staging receipt refers to another source: {receipt_path}")
    if profile_path.exists():
        existing=json_read(profile_path)
        if existing.get("profile_name")!=project or existing.get("family")!=FAMILY:
            raise RuntimeError(f"Existing profile belongs to another project/family: {profile_path}")
        bound=int(((existing.get('structure') or {}).get('m_line') or {}).get('boundary_after_col') or 0)
        if bound!=layout['wbs_col']:
            raise RuntimeError(f"Existing project profile boundary={bound}, inferred WBS col={layout['wbs_col']}; review needed")
    print("[STAGE] source:",safe_copy(source,staged,expected))
    print("[STAGE] working:",safe_copy(source,working,expected))
    if not profile_path.exists():
        from functions.cleaner.bootstrap import build_general_document_profile
        base=ROOT/"profiles"/"families"/"GENERAL_DOCUMENT.base.json"
        build_general_document_profile(project_id=project, workbook_path=staged,
             base_profile_path=base, output_path=profile_path,sheets=sheets,
             profile_overrides=make_overrides(layout))
        print(f"[PROFILE] Created PROPOSED {profile_path}")
    else:
        print(f"[PROFILE] Preserved existing {profile_path}")
    if not receipt_path.exists():
        receipt={"schema":"cleaner.general_document_staging/0.1","project":project,"family":FAMILY,
          "state":"PENDING_SOURCE_AND_PROFILE_REVIEW", "source_name":selected.get("name"),
          "source_original_path":selected.get("source"),"source_sha256":expected,
          "staged_source":str(staged),"working_copy":str(working),"profile_path":str(profile_path),
          "prepared_at_utc":datetime.now(timezone.utc).isoformat(),
          "notes":"Provisional geometric inference; no network original modified."}
        receipt_path.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf8")
    print("[RECEIPT]",receipt_path)
    runner=ROOT/"RUN_GENERAL_DOCUMENT_V02.py"
    if not runner.exists():
        print(f"[READY] Staged. Runner missing: {runner}")
        return 2
    # Resume rather than silently creating another run after a review gate. Validate
    # run identity and SHA before reusing the run's review/audit evidence.
    matching=[]
    for manifest in sorted((ROOT/"runs").glob(f"{project.lower()}_*/run_manifest.json")):
        try:
            run=json_read(manifest)
            entry=run.get("workbook") or {}
            if (str(run.get("profile_name", "")).upper()==project
                and Path(str(entry.get("path") or "")).resolve()==working.resolve()
                and str(entry.get("sha256_at_run_start") or "").lower()==expected):
                matching.append(manifest.parent.name)
        except (ValueError, OSError):
            continue
    if args.run_id:
        if args.run_id not in matching:
            raise RuntimeError(f"Requested --run-id {args.run_id} does not match staged workbook and source SHA")
        run_mode=["--run-id",args.run_id]
    elif len(matching)==1:
        run_mode=["--run-id",matching[0]]
        print(f"[RESUME] Existing matching run {matching[0]}")
    elif len(matching)>1:
        print(f"[REVIEW REQUIRED] Multiple runs match the same source: {', '.join(matching)}")
        print("  Rerun with --run-id to select an existing run; no new run created.")
        return 2
    else:
        run_mode=["--new-run"]
    print("[START] Review-driven Cleaner pipeline: human intervention only at review gates.", flush=True)
    cmd=[sys.executable,str(runner),"--project",project,*run_mode,"--workbook",str(working),"--status-date",args.status_date]
    env=dict(os.environ)
    for var in ("OPENBLAS_NUM_THREADS","OMP_NUM_THREADS","MKL_NUM_THREADS","NUMEXPR_NUM_THREADS"):
        env[var]="1"
    return subprocess.run(cmd,cwd=ROOT,env=env,check=False).returncode


if __name__=="__main__":
    try:
        raise SystemExit(main())
    except (ValueError,OSError,RuntimeError,KeyError) as exc:
        print(f"[STOP] {exc}",file=sys.stderr)
        raise SystemExit(1)
