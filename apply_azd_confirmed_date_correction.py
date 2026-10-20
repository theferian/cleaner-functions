"""Apply a human-confirmed AZD General Document date typo through Cleaner overlays.

Run from the Cleaner repository root, after copying this script there:
    python .\apply_azd_confirmed_date_correction.py --run-dir runs\azd_20261014T071802

This never modifies the source workbook. It persists an approved correction
in the run's resolved_state and project_solutions, with a guarded reference
and identity match to prevent correcting an unrelated correspondence event.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime
from pathlib import Path
import json
import shutil
import sys

from openpyxl.utils.cell import coordinate_from_string, range_boundaries

from cleaner_core.solutions import (
    load_project_solutions,
    save_project_solutions,
    solution_from_decision,
    upsert_solution,
)
from cleaner_core.state import ingest_review_decisions, load_resolved_state, save_resolved_state
from functions.cleaner.service import load_cleaner_run

SPEC = {
    "project_id": "AZD",
    "source_sheet": "230KV SS",
    "document_name": "PROT & CONTROL SINGLE LINE DIAGRAM",
    "document_wbs": "1.5.1.1.2.1.2",
    "event_side": "send",
    "letter_reference": "1404/ES/S/3003",
    "old_date": "1404/08/11",
    "corrected_date": "1404/07/19",
    "paired_receive_date": "1404/08/10",
    "paired_receive_status": "A",
    "confirmation": "User-confirmed source typo, 2026-10-16",
}
REVIEW_ID = "AZD_CONFIRMED_CORRESPONDENCE_SEND_DATE_1404_ES_S_3003"


def _normalized(value: object) -> str:
    return " ".join(str(value or "").split()).casefold()


def _find_event(run):
    """Resolve a letter reference to its date, not an unstable Excel row number."""
    sheets = run.document_regions().get("sheets") or {}
    card_sheets = run.card_maps().get("sheets") or {}
    sheet = SPEC["source_sheet"]
    regions = (sheets.get(sheet) or {}).get("regions") or []
    matches = [
        r for r in regions
        if str(((r.get("fields") or {}).get("wbs") or {}).get("value")) == SPEC["document_wbs"]
        and _normalized(((r.get("fields") or {}).get("name") or {}).get("value"))
        == _normalized(SPEC["document_name"])
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Document identity ambiguous or not found: {len(matches)} matches")
    region = matches[0]
    cards = [
        c for c in (card_sheets.get(sheet) or {}).get("cards") or []
        if int(c.get("record_index", -1)) == int(region["record_index"])
    ]
    if len(cards) != 1:
        raise RuntimeError(f"Corresponding card ambiguous or not found: {len(cards)} matches")
    card = cards[0]
    slots = {str(s.get("address")).upper(): s.get("value") for s in card.get("slots") or []}
    refs = [
        addr for addr, value in slots.items()
        if _normalized(value) == _normalized(SPEC["letter_reference"])
    ]
    if len(refs) != 1:
        raise RuntimeError(f"Letter reference ambiguous or not found: {len(refs)} matches")
    col, ref_row = coordinate_from_string(refs[0])
    if ref_row <= 1:
        raise RuntimeError("Letter reference cannot have its matching date above it")
    target = f"{col}{ref_row - 1}"
    observed = slots.get(target)
    if str(observed) != SPEC["old_date"]:
        if str(observed) == SPEC["corrected_date"]:
            raise RuntimeError("Source already contains the corrected date; no override required")
        raise RuntimeError(f"Unexpected old date at {sheet}!{target}: {observed!r}")
    from openpyxl.utils import column_index_from_string, get_column_letter
    next_col = get_column_letter(column_index_from_string(col) + 1)
    status_col = get_column_letter(column_index_from_string(col) + 2)
    def value_at(address: str):
        x1, y1, _, _ = range_boundaries(address)
        for slot in card.get("slots") or []:
            x_min, y_min, x_max, y_max = range_boundaries(slot.get("address"))
            if x_min <= x1 <= x_max and y_min <= y1 <= y_max:
                return slot.get("value")
        return None

    recv = value_at(f"{next_col}{ref_row - 1}")
    status = value_at(f"{status_col}{ref_row - 1}")
    if str(recv) != SPEC["paired_receive_date"] or str(status) != SPEC["paired_receive_status"]:
        raise RuntimeError(
            f"Correspondence does not match approved receive evidence: "
            f"{next_col}{ref_row-1}={recv!r}, {status_col}{ref_row-1}={status!r}"
        )
    return region, target


def _make_review_item(region: dict, target: str) -> dict:
    sheet = SPEC["source_sheet"]
    identity = {
        "name": SPEC["document_name"],
        "wbs": SPEC["document_wbs"],
    }
    return {
        "review_id": REVIEW_ID,
        "source": {
            "test_id": "CONFIRMED_CORRESPONDENCE_DATE",
            "status": "FAIL",
            "severity": "HIGH",
        },
        "location": {"sheet": sheet, "target": target},
        "observed": SPEC["old_date"],
        "message": "Human-confirmed incorrect send date for the specified correspondence letter",
        "evidence": {
            "document_identity": identity,
            "field_role": "send_date",
            "audit_kind": "RECEIVE_BEFORE_MATCHED_SEND",
            "event_side": SPEC["event_side"],
            "letter_reference": SPEC["letter_reference"],
            "paired_receive_date": SPEC["paired_receive_date"],
        },
        "decision": {
            "action": "SET_VALUE",
            "parameters": {
                "sheet": sheet,
                "target": target,
                "value": SPEC["corrected_date"],
                "expected_old_value": SPEC["old_date"],
                "field_role": "send_date",
            },
            "reason": SPEC["confirmation"] + "; original evidence preserved in raw card_map",
            "origin": "HUMAN_REVIEW",
            "decided_at": "2026-10-16T00:00:00",
        },
    }


def apply(run_dir: Path, *, dry_run: bool = False) -> dict:
    run = load_cleaner_run(run_dir)
    if run.profile_name.upper() != SPEC["project_id"]:
        raise RuntimeError(f"Refusing to modify non-AZD Cleaner run: {run.profile_name!r}")
    region, target = _find_event(run)
    sheet = SPEC["source_sheet"]

    state = load_resolved_state(run.state_path)
    override_key = f"{sheet}!{target}".upper()
    current_override = (state.get("value_overrides") or {}).get(override_key)
    if current_override is not None and current_override.get("value") != SPEC["corrected_date"]:
        raise RuntimeError(f"Conflicting reviewed correction already exists at {override_key}")

    path = run.project_solutions_path or (run.run_dir.parents[1] / "project_solutions" / "AZD.solutions.json")
    registry = load_project_solutions(path, project_id="AZD", create=True)
    item = _make_review_item(region, target)
    candidate = solution_from_decision(item, project_id="AZD", run_id=run.run_id)
    if candidate is None:
        raise RuntimeError("Cannot create project solution from reviewed correction")
    # Key the durable match on correspondence identity as well as document identity.
    # This custom test_id prevents unrelated DATE_CHRONOLOGY findings from matching.
    candidate["match"]["letter_reference"] = SPEC["letter_reference"]
    candidate["match"]["event_side"] = SPEC["event_side"]
    candidate["match"]["expected_old_date"] = SPEC["old_date"]
    import hashlib
    signature_fields = {k: v for k, v in candidate["match"].items() if k != "target_fallback"}
    candidate["signature"] = hashlib.sha1(json.dumps(signature_fields,sort_keys=True,ensure_ascii=False).encode("utf-8")).hexdigest()[:16]

    # Persist decision in standard Cleaner state and project solution store.
    payload = {"run": {"script_id": "CONFIRMED_CORRESPONDENCE_DATE", "run_id": run.run_id}, "decisions": [item]}
    ingest_review_decisions(state, payload)
    saved, created = upsert_solution(registry, candidate)

    evidence = {
        "schema": "cleaner.confirmed_correspondence_correction/0.1",
        "project": "AZD",
        "run_id": run.run_id,
        "source_sheet": sheet,
        "source_date_cell": target,
        "document_name": SPEC["document_name"],
        "document_wbs": SPEC["document_wbs"],
        "send_letter_reference": SPEC["letter_reference"],
        "old_date": SPEC["old_date"],
        "corrected_date": SPEC["corrected_date"],
        "paired_receive_date": SPEC["paired_receive_date"],
        "paired_receive_status": SPEC["paired_receive_status"],
        "review_id": REVIEW_ID,
        "project_solution_handle": saved["handle"],
        "solution_created": created,
        "confirmation": SPEC["confirmation"],
        "source_workbook_mutated": False,
    }
    if dry_run:
        return evidence

    # Back up the two mutable JSON files before writing anything.
    backup = run.run_dir / "confirmed_corrections" / "backup"
    backup.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    for src in (run.state_path, path):
        if src.exists():
            shutil.copy2(src, backup / f"{src.stem}.{timestamp}.json")

    save_resolved_state(state, run.state_path)
    save_project_solutions(registry, path)
    evidence_path = run.run_dir / "confirmed_corrections" / "AZD_PROT_CONTROL_SEND_DATE_1404_ES_S_3003.json"
    evidence_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    # Append to the existing work list without disturbing other project tasks.
    tasks_path = run.run_dir.parents[1] / "docs" / "DEMO_TASK_LIST.md"
    checklist = (
        "- [x] 4A.13a Record approved AZD send-date correction "
        "(PROT & CONTROL, letter 1404/ES/S/3003).",
        "- [ ] 4A.13b Add guarded project-solution replay to every new Cleaner run.",
        "- [ ] 4A.13c Add generic date-contradiction review → approved repair → validation workflow.",
        "- [ ] 4A.13d Revalidate AZD chronology diagnostics and regenerate downstream reports.",
    )
    tasks_path.parent.mkdir(parents=True, exist_ok=True)
    original = tasks_path.read_text(encoding="utf-8") if tasks_path.exists() else "# Cleaner Worklist\n"
    missing = [line for line in checklist if line not in original]
    if missing:
        tasks_path.write_text(original.rstrip() + "\n\n## Confirmed source-date corrections\n" + "\n".join(missing) + "\n", encoding="utf-8")
    return evidence


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--run-dir", required=True, type=Path)
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()
    result = apply(args.run_dir, dry_run=args.dry_run)
    print("[DRY RUN]" if args.dry_run else "[APPLIED]", result["source_sheet"], result["source_date_cell"],
          result["old_date"], "->", result["corrected_date"],
          "solution=", result["project_solution_handle"])
    print("[NEXT] Re-run board-to-log, coverage audit, then engineering-lists for this run")


if __name__ == "__main__":
    main()
