from __future__ import annotations

from dataclasses import dataclass, field, asdict
from typing import Any

from openpyxl.utils import get_column_letter

from .document_regions import DocumentRegion, scan_document_regions
from .ranges import Rect


@dataclass(frozen=True)
class CardSlot:
    address: str
    row_offset_start: int
    row_offset_end: int
    col_start: int
    col_end: int
    merged: bool
    value: Any = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class CardMap:
    sheet: str
    record_index: int
    anchor_row: int
    start_row: int
    end_row: int
    identifier_range: str
    data_range: str
    data_start_col: int
    data_end_col: int
    slots: tuple[CardSlot, ...] = ()

    @property
    def target(self) -> str:
        return f"{self.identifier_range}|{self.data_range}"

    def to_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["target"] = self.target
        payload["slot_count"] = len(self.slots)
        return payload


@dataclass
class CardMapScan:
    sheet: str
    cards: list[CardMap] = field(default_factory=list)
    issues: list[dict[str, Any]] = field(default_factory=list)
    data_start_col: int | None = None
    data_end_col: int | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "data_start_col": self.data_start_col,
            "data_end_col": self.data_end_col,
            "cards": [x.to_dict() for x in self.cards],
            "issues": list(self.issues),
            "summary": {
                "card_count": len(self.cards),
                "cards_with_data": sum(1 for x in self.cards if x.slots),
                "cards_without_data": sum(1 for x in self.cards if not x.slots),
                "issue_count": len(self.issues),
            },
        }


def card_map_config(profile: dict) -> dict[str, Any]:
    configured = profile.get("data_model", {}).get("card_map", {})
    boundary_after_col = int(
        profile.get("structure", {}).get("m_line", {}).get("boundary_after_col", 3)
    )
    return {
        "enabled": configured.get("enabled", True),
        "data_start_col": int(configured.get("data_start_col", boundary_after_col + 1)),
        "data_end_strategy": str(configured.get("data_end_strategy", "MERGE_EXTENT")).upper(),
        "semantic_fields": list(configured.get("semantic_fields", [
            "reference", "send_date", "receive_date", "status"
        ])),
        "field_map": dict(configured.get("field_map", {})),
    }


def _intersects_rows(rect: Rect, start_row: int, end_row: int) -> bool:
    return not (rect.r2 < start_row or rect.r1 > end_row)


def _sheet_data_end_col(snapshot, sheet: str, regions: list[DocumentRegion], data_start_col: int) -> int:
    """Find the structural right edge from merge geometry, with worksheet max-column fallback."""
    _, max_col = snapshot.dimensions(sheet)
    merged = snapshot.merged_ranges(sheet)
    if not regions:
        return max(data_start_col, max_col)
    first_row = min(x.start_row for x in regions)
    last_row = max(x.end_row for x in regions)
    structural = [
        rect.c2 for rect in merged
        if rect.c2 >= data_start_col and _intersects_rows(rect, first_row, last_row)
    ]
    if structural:
        return max(data_start_col, max(structural))
    return max(data_start_col, max_col)


def _address(c1: int, r1: int, c2: int, r2: int) -> str:
    a = f"{get_column_letter(c1)}{r1}"
    b = f"{get_column_letter(c2)}{r2}"
    return a if a == b else f"{a}:{b}"


def _slots_for_region(snapshot, sheet: str, region: DocumentRegion, data_start_col: int, data_end_col: int) -> tuple[CardSlot, ...]:
    slots: list[CardSlot] = []
    covered: set[tuple[int, int]] = set()

    for rect in snapshot.merged_ranges(sheet):
        if rect.r1 < region.start_row or rect.r2 > region.end_row:
            continue
        if rect.c2 < data_start_col or rect.c1 > data_end_col:
            continue
        c1 = max(rect.c1, data_start_col)
        c2 = min(rect.c2, data_end_col)
        value = snapshot.value(sheet, f"{get_column_letter(rect.c1)}{rect.r1}")
        slots.append(CardSlot(
            address=_address(c1, rect.r1, c2, rect.r2),
            row_offset_start=rect.r1 - region.start_row,
            row_offset_end=rect.r2 - region.start_row,
            col_start=c1,
            col_end=c2,
            merged=True,
            value=value,
        ))
        for r in range(rect.r1, rect.r2 + 1):
            for c in range(c1, c2 + 1):
                covered.add((r, c))

    # Keep unmerged, populated cells visible in the map too. Empty grid cells are noise.
    for r in range(region.start_row, region.end_row + 1):
        for c in range(data_start_col, data_end_col + 1):
            if (r, c) in covered:
                continue
            addr = f"{get_column_letter(c)}{r}"
            value = snapshot.value(sheet, addr)
            if value is None or str(value).strip() == "":
                continue
            slots.append(CardSlot(
                address=addr,
                row_offset_start=r - region.start_row,
                row_offset_end=r - region.start_row,
                col_start=c,
                col_end=c,
                merged=False,
                value=value,
            ))

    return tuple(sorted(slots, key=lambda x: (x.row_offset_start, x.col_start, x.row_offset_end, x.col_end)))


def scan_card_map(snapshot, profile: dict, sheet: str) -> CardMapScan:
    config = card_map_config(profile)
    scan = CardMapScan(sheet=sheet)
    if not snapshot.has_sheet(sheet):
        scan.issues.append({"kind": "SHEET_NOT_FOUND", "message": f"{sheet}: sheet not found."})
        return scan

    boundary_scan = scan_document_regions(snapshot, profile, sheet)
    regions = boundary_scan.regions
    data_start_col = config["data_start_col"]
    data_end_col = _sheet_data_end_col(snapshot, sheet, regions, data_start_col)
    scan.data_start_col = data_start_col
    scan.data_end_col = data_end_col

    boundary_after_col = int(profile.get("structure", {}).get("m_line", {}).get("boundary_after_col", data_start_col - 1))

    if data_end_col < data_start_col:
        scan.issues.append({
            "kind": "NO_DATA_AREA",
            "message": f"{sheet}: data area ends before it starts.",
            "data_start_col": data_start_col,
            "data_end_col": data_end_col,
        })
        return scan

    for region in regions:
        identifier_range = _address(1, region.start_row, boundary_after_col, region.end_row)
        data_range = _address(data_start_col, region.start_row, data_end_col, region.end_row)
        slots = _slots_for_region(snapshot, sheet, region, data_start_col, data_end_col)
        scan.cards.append(CardMap(
            sheet=sheet,
            record_index=region.record_index,
            anchor_row=region.anchor_row,
            start_row=region.start_row,
            end_row=region.end_row,
            identifier_range=identifier_range,
            data_range=data_range,
            data_start_col=data_start_col,
            data_end_col=data_end_col,
            slots=slots,
        ))

    if regions and not any(card.slots for card in scan.cards):
        scan.issues.append({
            "kind": "NO_DATA_SLOTS",
            "message": f"{sheet}: document cards were found but no populated/merged data slots were mapped.",
        })

    return scan


def scan_all_card_maps(snapshot, profile: dict) -> dict[str, CardMapScan]:
    target_sheets = (
        profile.get("structure", {}).get("m_line", {}).get("target_sheets")
        or snapshot.sheet_names()
    )
    return {sheet: scan_card_map(snapshot, profile, sheet) for sheet in target_sheets}
