from __future__ import annotations

import re
from typing import Any

from .document_regions import scan_all_document_regions
from .ranges import parse_range
from openpyxl.utils import get_column_letter

_VALID_WBS = re.compile(r"^\d+(?:\.\d+)*$")


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def document_identity(region) -> dict[str, Any]:
    """Return a row-independent identity for one logical document region.

    Name is the primary project-stable anchor. A valid WBS is included when
    present, but placeholder/invalid WBS values are deliberately omitted so a
    WBS repair can still be rediscovered after the document moves rows.
    """
    identity: dict[str, Any] = {
        "record_index": int(region.record_index),
    }
    durable_keys = (
        "document_type", "document_type_a", "document_type_b",
        "category", "type", "name", "code", "wbs",
    )
    for key in durable_keys:
        anchor = region.field(key)
        value = _clean(anchor.value if anchor is not None else None)
        if value is None:
            continue
        if key == "wbs" and _VALID_WBS.fullmatch(value) is None:
            continue
        if key == "code" and value in {"-", "–", "—"}:
            continue
        identity[key] = value
    return identity


def build_document_context_index(snapshot, profile: dict) -> dict[str, list[dict[str, Any]]]:
    scans = scan_all_document_regions(snapshot, profile)
    index: dict[str, list[dict[str, Any]]] = {}
    for sheet, scan in scans.items():
        index[sheet] = [
            {
                "start_row": int(region.start_row),
                "end_row": int(region.end_row),
                "identity": document_identity(region),
            }
            for region in scan.regions
        ]
    return index




def document_relative_target(target: str | None, region_start_row: int) -> str | None:
    """Return a row-shift-resistant target within a logical document region.

    Examples:
      H11      in a region starting at row 11 -> H+0
      H11:H12 in a region starting at row 11 -> H+0:H+1

    Absolute row numbers are deliberately removed while columns and vertical
    offsets within the document card are preserved.
    """
    if not target:
        return None
    try:
        rect = parse_range(str(target))
    except Exception:
        return None
    left = f"{get_column_letter(rect.c1)}+{rect.r1 - int(region_start_row)}"
    right = f"{get_column_letter(rect.c2)}+{rect.r2 - int(region_start_row)}"
    return left if left == right else f"{left}:{right}"


def _target_row(target: str | None) -> int | None:
    if not target:
        return None
    try:
        return int(parse_range(str(target)).r1)
    except Exception:
        return None


def attach_document_context(findings: list, snapshot, profile: dict) -> list:
    """Attach stable document identity to findings before review/solution matching.

    The function mutates Finding.evidence in place and returns the same list.
    Findings outside logical document regions (e.g. separator geometry) are
    intentionally left untouched.
    """
    index = build_document_context_index(snapshot, profile)
    for finding in findings:
        sheet = getattr(finding, "sheet", None)
        row = _target_row(getattr(finding, "target", None))
        if not sheet or row is None:
            continue
        for entry in index.get(sheet, []):
            if entry["start_row"] <= row <= entry["end_row"]:
                evidence = getattr(finding, "evidence", None)
                if evidence is None:
                    evidence = {}
                    finding.evidence = evidence
                evidence.setdefault("document_identity", dict(entry["identity"]))
                relative = document_relative_target(
                    getattr(finding, "target", None),
                    int(entry["start_row"]),
                )
                if relative:
                    evidence.setdefault("document_relative_target", relative)
                break
    return findings
