"""Safely replay human-approved correspondence-date corrections in new Cleaner runs.

Approved project-solution entries retain the user's decision. Source evidence from
an earlier confirmed correction supplies guards missing from the original solution
record. A mismatch never triggers a guessed repair or an Excel write.
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any
import json

from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

from .state import ingest_review_decisions


SCHEMA = "cleaner.correspondence_replay/0.1"
TEST_ID = "CONFIRMED_CORRESPONDENCE_DATE"


def _norm(value: Any) -> str:
    return " ".join(str(value or "").split()).casefold()


def _json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _value_at(slots: list[dict], target: str) -> Any:
    x, y, _, _ = range_boundaries(target)
    found = [s.get("value") for s in slots
             if _contains(s.get("address"), x, y)]
    if len(found) > 1:
        raise ValueError("Overlapping ambiguous slot coordinates: " + target)
    return found[0] if found else None


def _contains(address: Any, col: int, row: int) -> bool:
    if not address:
        return False
    x1, y1, x2, y2 = range_boundaries(str(address))
    return x1 <= col <= x2 and y1 <= row <= y2


def _lookup_guard(solution: dict, runs_root: Path) -> dict | None:
    """Read prior human-confirmed evidence; never reconstruct missing approval."""
    attached = solution.get("correspondence_guard")
    if isinstance(attached, dict):
        return attached
    source = solution.get("source") or {}
    old_run = str(source.get("run_id") or "")
    if not old_run or Path(old_run).name != old_run:
        return None
    evidence_dir = runs_root / old_run / "confirmed_corrections"
    for path in sorted(evidence_dir.glob("*.json")):
        evidence = _json(path)
        if (str(evidence.get("project_solution_handle")) == str(solution.get("handle"))
                and _norm(evidence.get("send_letter_reference")) == _norm((solution.get("match") or {}).get("letter_reference"))):
            return evidence
    return None


def _apply_one(solution: dict, guard: dict, regions: dict, cards: dict, state: dict, run_id: str) -> dict:
    match = solution.get("match") or {}
    identity = match.get("document_identity") or {}
    sheet = match.get("sheet")
    ref = match.get("letter_reference")
    old = guard.get("old_date")
    new = (solution.get("resolution") or {}).get("parameters", {}).get("value")
    receive_date = guard.get("paired_receive_date")
    receive_status = guard.get("paired_receive_status")
    old_guard = match.get("expected_old_date") or match.get("observed")

    required = [sheet, ref, identity.get("name"), identity.get("wbs"), old, new,
                receive_date, receive_status, old_guard]
    if any(x in (None, "") for x in required) or str(old) != str(old_guard):
        return {"result": "NEEDS_REVIEW", "reason": "INCOMPLETE_OR_CONFLICTING_GUARDS"}
    if _norm(guard.get("document_name")) != _norm(identity.get("name")) or str(guard.get("document_wbs")) != str(identity.get("wbs")):
        return {"result": "NEEDS_REVIEW", "reason": "GUARD_DOCUMENT_MISMATCH"}
    if str(guard.get("source_sheet")) != str(sheet):
        return {"result": "NEEDS_REVIEW", "reason": "GUARD_SHEET_MISMATCH"}

    all_regions = (regions.get("sheets") or {}).get(sheet, {}).get("regions", [])
    located = [r for r in all_regions
               if _norm(((r.get("fields") or {}).get("name") or {}).get("value")) == _norm(identity["name"])
               and str(((r.get("fields") or {}).get("wbs") or {}).get("value")) == str(identity["wbs"])]
    if len(located) != 1:
        return {"result": "NEEDS_REVIEW", "reason": "DOCUMENT_NOT_UNIQUE", "matches": len(located)}
    region = located[0]
    all_cards = (cards.get("sheets") or {}).get(sheet, {}).get("cards", [])
    located_cards = [c for c in all_cards if int(c.get("record_index", -1)) == int(region["record_index"])]
    if len(located_cards) != 1:
        return {"result": "NEEDS_REVIEW", "reason": "CARD_NOT_UNIQUE", "matches": len(located_cards)}
    slots = located_cards[0].get("slots") or []
    candidates = [s for s in slots if _norm(s.get("value")) == _norm(ref)]
    if len(candidates) != 1:
        return {"result": "NEEDS_REVIEW", "reason": "LETTER_REFERENCE_NOT_UNIQUE", "matches": len(candidates)}

    try:
        address = str(candidates[0]["address"])
        col_text, row = coordinate_from_string(address)
        if row < 2:
            raise ValueError("letter reference cannot be on the first row")
        col = column_index_from_string(col_text)
        date_target = f"{col_text}{row - 1}"
        received_at = f"{get_column_letter(col + 1)}{row - 1}"
        status_at = f"{get_column_letter(col + 2)}{row - 1}"
        observed = _value_at(slots, date_target)
        actual_receive = _value_at(slots, received_at)
        actual_status = _value_at(slots, status_at)
    except (ValueError, TypeError, KeyError) as exc:
        return {"result": "NEEDS_REVIEW", "reason": "INVALID_SLOT_GEOMETRY", "detail": str(exc)}

    info = {"sheet": sheet, "target": date_target, "observed": observed,
            "expected_old": old, "approved_value": new, "reference": ref}
    if _norm(actual_receive) != _norm(receive_date) or _norm(actual_status) != _norm(receive_status):
        return {"result": "NEEDS_REVIEW", "reason": "PAIRED_RECEIVE_GUARD_FAILED", **info,
                "actual_receive_date": actual_receive, "actual_receive_status": actual_status}
    state_key = f"{sheet}!{date_target}".upper()
    prior_override = (state.get("value_overrides") or {}).get(state_key)
    if prior_override:
        if str(prior_override.get("value")) == str(new):
            return {"result": "ALREADY_APPLIED", **info}
        return {"result": "NEEDS_REVIEW", "reason": "EXISTING_OVERRIDE_CONFLICT", **info}
    if str(observed) == str(new):
        return {"result": "ALREADY_CORRECT", **info}
    if str(observed) != str(old):
        return {"result": "NEEDS_REVIEW", "reason": "OLD_VALUE_GUARD_FAILED", **info}

    review_id = f"REPLAY_{run_id}_{solution.get('handle')}_{sheet}_{date_target}".replace(" ", "_")
    item = {"review_id": review_id, "source": {"test_id": TEST_ID, "status": "REVIEWED"},
            "location": {"sheet": sheet, "target": date_target}, "observed": observed,
            "evidence": {"document_identity": deepcopy(identity), "letter_reference": ref,
                         "paired_receive_date": receive_date, "paired_receive_status": receive_status},
            "decision": {"action": "SET_VALUE", "origin": "PROJECT_SOLUTION",
                         "project_solution_handle": solution.get("handle"),
                         "reason": f"Replayed approved correspondence correction {solution.get('handle')}",
                         "parameters": {"sheet": sheet, "target": date_target,
                                        "value": new, "expected_old_value": old, "field_role": "send_date"}}}
    ingest_review_decisions(state, {"run": {"script_id": TEST_ID, "run_id": run_id}, "decisions": [item]})
    return {"result": "APPLIED", **info}


def replay_approved_correspondence_corrections(
        run_dir: str | Path, project_solutions: dict, state: dict, *, dry_run: bool = False) -> dict:
    """Apply guarded, approved corrections to a new run's resolved-state only.

    Requires completed DOCUMENT_BOUNDARIES/CARD_MAP stages. Source workbook is
    never modified. Returns evidence for a per-run report and review gate.
    """
    run_dir = Path(run_dir)
    records = [s for s in (project_solutions.get("solutions") or [])
               if s.get("enabled", True)
               and _norm((s.get("match") or {}).get("test_id")) == _norm(TEST_ID)]
    output = {"schema": SCHEMA, "run_id": run_dir.name, "results": [], "summary": {}}
    if not records:
        output["summary"] = {"APPLIED": 0, "NEEDS_REVIEW": 0}
        return output
    region_file = run_dir / "document_regions.json"
    cards_file = run_dir / "card_maps.json"
    if not region_file.exists() or not cards_file.exists():
        raise RuntimeError("Correspondence replay requires document_regions.json and card_maps.json")
    regions, cards = _json(region_file), _json(cards_file)
    working = deepcopy(state) if dry_run else state
    for solution in records:
        guard = _lookup_guard(solution, run_dir.parent)
        if guard is None:
            result = {"result": "NEEDS_REVIEW", "reason": "APPROVAL_EVIDENCE_UNAVAILABLE"}
        else:
            result = _apply_one(solution, guard, regions, cards, working, run_dir.name)
        result["solution_handle"] = solution.get("handle")
        output["results"].append(result)
    output["summary"] = {key: sum(x["result"] == key for x in output["results"])
                         for key in ("APPLIED", "ALREADY_APPLIED", "ALREADY_CORRECT", "NEEDS_REVIEW")}
    return output


def _main() -> int:
    import argparse
    from .state import load_resolved_state
    from .solutions import load_project_solutions

    parser = argparse.ArgumentParser(description="Preview approved correspondence replay (read-only)")
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--project", required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    state = load_resolved_state(args.run_dir / "resolved_state.json")
    registry = load_project_solutions(
        root / "project_solutions" / f"{args.project}.solutions.json",
        project_id=args.project,
    )
    report = replay_approved_correspondence_corrections(
        args.run_dir, registry, state, dry_run=True
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not report["summary"].get("NEEDS_REVIEW", 0) else 2


if __name__ == "__main__":
    raise SystemExit(_main())
