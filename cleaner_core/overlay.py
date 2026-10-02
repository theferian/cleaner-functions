from __future__ import annotations

from copy import deepcopy
from openpyxl.utils import get_column_letter

from .ranges import Rect, parse_range


def _key(sheet: str, target: str) -> str:
    return f"{sheet}!{target}".upper()


class ResolvedWorkbookView:
    """
    Read-only logical view of a workbook snapshot plus reviewed resolutions.

    Later tests read this object instead of the raw snapshot. Nothing is
    physically written to the user's Excel workbook.
    """

    def __init__(self, snapshot, state: dict):
        self.base = snapshot
        self.state = state

    @property
    def path(self):
        return getattr(self.base, "path", None)

    def close(self):
        return self.base.close()

    def sheet_names(self):
        return self.base.sheet_names()

    def has_sheet(self, sheet: str):
        return self.base.has_sheet(sheet)

    def dimensions(self, sheet: str):
        return self.base.dimensions(sheet)

    def _value_override(self, sheet: str, address: str):
        return self.state.get("value_overrides", {}).get(_key(sheet, address))

    def value(self, sheet: str, address: str):
        override = self._value_override(sheet, address)
        if override is not None:
            return override.get("value")
        return self.base.value(sheet, address)

    def effective_value(self, sheet: str, address: str):
        info = self.merge_info(sheet, address)
        return info["value"]

    def merged_ranges(self, sheet: str) -> list[Rect]:
        ranges = list(self.base.merged_ranges(sheet))

        for op in self.state.get("merge_operations", []):
            if str(op.get("sheet") or "") != sheet:
                continue
            target = str(op.get("target") or "")
            if not target:
                continue

            action = str(op.get("action") or "").upper()
            target_rect = parse_range(target)

            if action == "UNMERGE_RANGE":
                ranges = [
                    rect for rect in ranges
                    if rect.address.upper() != target_rect.address.upper()
                ]
            elif action == "MERGE_RANGE":
                if not any(
                    rect.address.upper() == target_rect.address.upper()
                    for rect in ranges
                ):
                    ranges.append(target_rect)

        return ranges

    def merged_range_containing(self, sheet: str, address: str) -> Rect | None:
        cell = parse_range(address)
        row, col = cell.r1, cell.c1
        for rect in self.merged_ranges(sheet):
            if rect.r1 <= row <= rect.r2 and rect.c1 <= col <= rect.c2:
                return rect
        return None

    def merge_info(self, sheet: str, address: str):
        rect = self.merged_range_containing(sheet, address)
        if rect is None:
            return {
                "merged": False,
                "range": None,
                "anchor_row": None,
                "anchor_col": None,
                "value": self.value(sheet, address),
            }

        anchor = f"{get_column_letter(rect.c1)}{rect.r1}"
        return {
            "merged": True,
            "range": rect.address,
            "anchor_row": rect.r1,
            "anchor_col": rect.c1,
            "value": self.value(sheet, anchor),
        }

    def anchor_value(self, sheet: str, rect: Rect):
        return self.value(sheet, f"{get_column_letter(rect.c1)}{rect.r1}")


def apply_state_to_profile(profile: dict, state: dict) -> dict:
    """
    Return a profile view containing accepted exceptions from prior review.

    Free-text PROFILE_CHANGE decisions are intentionally *not* auto-applied;
    they remain state/audit items until we define a structured profile-patch
    contract.
    """
    resolved = deepcopy(profile)
    known = resolved.setdefault("known_exceptions", [])
    existing_ids = {str(x.get("id")) for x in known}

    for exc in state.get("accepted_exceptions", []):
        review_id = str(exc.get("source_review_id") or "")
        rule_id = f"RESOLVED_{review_id}" if review_id else "RESOLVED_EXCEPTION"
        if rule_id in existing_ids:
            continue

        match = {
            "test_id": exc.get("test_id"),
            "sheet": exc.get("sheet") or "*",
        }
        if exc.get("rows"):
            match["rows"] = exc.get("rows")
        if exc.get("target"):
            match["target"] = exc.get("target")

        known.append({
            "id": rule_id,
            "enabled": True,
            "reason": exc.get("reason") or "Accepted through resolved-state review.",
            "match": match,
            "source": {
                "kind": "RESOLVED_STATE",
                "review_id": review_id,
            },
        })
        existing_ids.add(rule_id)

    return resolved


def build_resolved_context(snapshot, profile: dict, state: dict):
    return ResolvedWorkbookView(snapshot, state), apply_state_to_profile(profile, state)
