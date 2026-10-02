from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .ranges import Rect, parse_range


class WorkbookSnapshot:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        keep_vba = self.path.suffix.lower() in {".xlsm", ".xltm"}
        self._wb = load_workbook(
            self.path,
            data_only=False,
            read_only=False,
            keep_vba=keep_vba,
        )

    def close(self):
        self._wb.close()

    def sheet_names(self) -> list[str]:
        return list(self._wb.sheetnames)

    def has_sheet(self, sheet: str) -> bool:
        return sheet in self._wb.sheetnames

    def value(self, sheet: str, address: str):
        return self._wb[sheet][address].value

    def dimensions(self, sheet: str) -> tuple[int, int]:
        ws = self._wb[sheet]
        return ws.max_row, ws.max_column

    def merged_ranges(self, sheet: str) -> list[Rect]:
        ws = self._wb[sheet]
        return [parse_range(str(rng)) for rng in ws.merged_cells.ranges]

    def merged_range_containing(self, sheet: str, address: str) -> Rect | None:
        cell = parse_range(address)
        row, col = cell.r1, cell.c1
        for rect in self.merged_ranges(sheet):
            if rect.r1 <= row <= rect.r2 and rect.c1 <= col <= rect.c2:
                return rect
        return None

    def anchor_value(self, sheet: str, rect: Rect):
        return self.value(sheet, f"{get_column_letter(rect.c1)}{rect.r1}")


def load_snapshot(path: str | Path) -> WorkbookSnapshot:
    return WorkbookSnapshot(path)
