from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass, asdict
from datetime import datetime
import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class SheetProbe:
    sheet: str
    max_row: int
    max_col: int
    identifier_rows: int
    transaction_cells: int
    header_hits: tuple[str, ...]
    selected: bool
    reason: str

    def to_dict(self) -> dict[str, Any]:
        payload = asdict(self)
        payload["header_hits"] = list(self.header_hits)
        return payload


HEADER_TERMS = (
    "send",
    "receive",
    "status",
    "rev",
    "ارسال",
    "دریافت",
    "وضعیت",
)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def inspect_general_document_workbook(path: str | Path) -> list[SheetProbe]:
    """Conservatively identify General-Document data sheets.

    A sheet is selected when it has repeated identifier content in B/C and
    non-empty transaction-board content at D or later.  Header terms are
    evidence only; they are deliberately not mandatory because historical
    project templates are inconsistent.
    """
    from openpyxl import load_workbook

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    keep_vba = path.suffix.lower() in {".xlsm", ".xltm"}
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=keep_vba)
    try:
        probes: list[SheetProbe] = []
        for ws in wb.worksheets:
            max_row = int(ws.max_row or 0)
            max_col = int(ws.max_column or 0)
            identifier_rows = 0
            transaction_cells = 0
            for row in range(3, max_row + 1):
                b = _text(ws.cell(row=row, column=2).value) if max_col >= 2 else ""
                c = _text(ws.cell(row=row, column=3).value) if max_col >= 3 else ""
                if b or c:
                    identifier_rows += 1
            if max_col >= 4:
                for row in range(1, min(max_row, 80) + 1):
                    for col in range(4, max_col + 1):
                        if _text(ws.cell(row=row, column=col).value):
                            transaction_cells += 1

            header_blob = " ".join(
                _text(ws.cell(row=r, column=c).value).casefold()
                for r in range(1, min(max_row, 3) + 1)
                for c in range(1, min(max_col, 40) + 1)
            )
            hits = tuple(term for term in HEADER_TERMS if term.casefold() in header_blob)
            selected = (
                ws.sheet_state == "visible"
                and max_col >= 4
                and identifier_rows >= 2
                and transaction_cells >= 2
            )
            if selected:
                reason = "identifier columns B/C + transaction board D+ detected"
            elif ws.sheet_state != "visible":
                reason = f"sheet_state={ws.sheet_state}"
            elif max_col < 4:
                reason = "fewer than four used columns"
            elif identifier_rows < 2:
                reason = "insufficient identifier rows in B/C"
            else:
                reason = "insufficient transaction-board content in D+"
            probes.append(
                SheetProbe(
                    sheet=str(ws.title),
                    max_row=max_row,
                    max_col=max_col,
                    identifier_rows=identifier_rows,
                    transaction_cells=transaction_cells,
                    header_hits=hits,
                    selected=selected,
                    reason=reason,
                )
            )
        return probes
    finally:
        wb.close()


def _sheet_instance() -> dict[str, Any]:
    return {
        "instance_type": "SHEET",
        "field_map": {
            "document_type": {"kind": "COLUMN", "column": 1},
            "name": {"kind": "COLUMN", "column": 2},
            "wbs": {"kind": "COLUMN", "column": 3},
        },
    }


def _identifier_field_from_layout(profile: dict[str, Any], spec: dict[str, Any]) -> dict[str, Any]:
    key = str(spec.get("key") or "").strip()
    if not key:
        raise ValueError("Identifier layout field is missing key.")
    existing = next(
        (deepcopy(x) for x in profile.get("data_model", {}).get("family_fields", []) if x.get("key") == key),
        None,
    )
    field = existing or {
        "key": key,
        "label": str(spec.get("label") or key.replace("_", " ").title()),
        "data_type": str(spec.get("data_type") or "TEXT"),
        "required": bool(spec.get("required", False)),
        "validation": deepcopy(spec.get("validation") or {"kind": "NONE"}),
        "cleaning": deepcopy(spec.get("cleaning") or {}),
    }
    field["label"] = str(spec.get("label") or field.get("label") or key)
    field["data_type"] = str(spec.get("data_type") or field.get("data_type") or "TEXT")
    if "required" in spec:
        field["required"] = bool(spec.get("required"))
    if "validation" in spec:
        field["validation"] = deepcopy(spec.get("validation") or {})
    if "cleaning" in spec:
        field["cleaning"] = deepcopy(spec.get("cleaning") or {})
    column = int(spec["column"])
    field["default_locator"] = {"kind": "COLUMN", "column": column}
    return field


def apply_general_document_profile_overrides(
    profile: dict[str, Any],
    *,
    selected_sheets: list[str],
    overrides: dict[str, Any] | None,
) -> dict[str, Any]:
    """Apply project-specific structural layout without importing project knowledge.

    This is intended for genuine template dialects such as FSH where the
    identifier side is Category/Type/Name/Code/WBS (A:E) rather than the
    family default three-column A:C layout.  Learned fixes and exceptions are
    deliberately outside this function.
    """
    overrides = dict(overrides or {})
    layout = dict(overrides.get("identifier_layout") or {})
    if not layout:
        profile.setdefault("data_model", {})["instances"] = {
            sheet: _sheet_instance() for sheet in selected_sheets
        }
        return profile

    specs = list(layout.get("fields") or [])
    if not specs:
        raise ValueError("identifier_layout.fields cannot be empty.")
    fields = [_identifier_field_from_layout(profile, dict(spec)) for spec in specs]
    columns = {str(x["key"]): int(x["default_locator"]["column"]) for x in fields}
    if "name" not in columns or "wbs" not in columns:
        raise ValueError("Project identifier layout must define at least name and wbs fields.")

    data_model = profile.setdefault("data_model", {})
    data_model["family_fields"] = fields
    data_model["instances"] = {
        sheet: {
            "instance_type": "SHEET",
            "field_map": {
                key: {"kind": "COLUMN", "column": column}
                for key, column in columns.items()
            },
        }
        for sheet in selected_sheets
    }

    signal_fields = list(layout.get("signal_fields") or columns.keys())
    boundary_fields = list(layout.get("boundary_fields") or ["name", "wbs"])
    separator_fields = list(layout.get("separator_fields") or boundary_fields)
    boundary_after = int(layout.get("boundary_after_col") or max(columns.values()))
    data_start = int(layout.get("data_start_col") or (boundary_after + 1))

    m_line = profile.setdefault("structure", {}).setdefault("m_line", {})
    m_line["boundary_after_col"] = boundary_after
    m_line["boundary_label"] = str(
        layout.get("boundary_label")
        or f"{get_column_letter(boundary_after)}|{get_column_letter(boundary_after + 1)}"
    )

    boundaries = data_model.setdefault("document_boundaries", {})
    boundaries["signal_fields"] = signal_fields
    boundaries["boundary_fields"] = boundary_fields
    boundaries["separator_fields"] = separator_fields

    detection = data_model.setdefault("record_detection", {})
    detection["signal_fields"] = signal_fields

    structure = data_model.setdefault("document_structure", {})
    structure["transaction_start_col"] = data_start

    card_map = data_model.setdefault("card_map", {})
    card_map["data_start_col"] = data_start

    return profile


def build_general_document_profile(
    *,
    project_id: str,
    workbook_path: str | Path,
    base_profile_path: str | Path,
    output_path: str | Path,
    sheets: Iterable[str] | None = None,
    overwrite: bool = False,
    profile_overrides: dict[str, Any] | None = None,
) -> tuple[Path, dict[str, Any]]:
    project_id = str(project_id).strip().upper()
    workbook_path = Path(workbook_path)
    base_profile_path = Path(base_profile_path)
    output_path = Path(output_path)

    if output_path.exists() and not overwrite:
        raise FileExistsError(
            f"Project profile already exists: {output_path}. Use overwrite=True only if you intend to rebuild it."
        )

    base = json.loads(base_profile_path.read_text(encoding="utf-8"))
    probes = inspect_general_document_workbook(workbook_path)
    if sheets is None:
        selected = [p.sheet for p in probes if p.selected]
    else:
        requested = [str(x).strip() for x in sheets if str(x).strip()]
        existing = {p.sheet for p in probes}
        missing = [x for x in requested if x not in existing]
        if missing:
            raise ValueError(f"Requested sheet(s) not found in workbook: {missing}")
        selected = requested

    if not selected:
        raise RuntimeError(
            "No General-Document data sheets were selected. Inspect the workbook probe and rerun with explicit --sheet values."
        )

    profile = deepcopy(base)
    profile["profile_name"] = project_id
    profile["workbook"] = {"expected_sheets": list(selected)}
    profile.setdefault("structure", {}).setdefault("m_line", {})["target_sheets"] = list(selected)
    profile = apply_general_document_profile_overrides(
        profile,
        selected_sheets=list(selected),
        overrides=profile_overrides,
    )
    profile["known_exceptions"] = []
    profile["known_fixes"] = []
    profile["bootstrap"] = {
        "template": "GENERAL_DOCUMENT",
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "project_id": project_id,
        "source_workbook": str(workbook_path),
        "selected_sheets": list(selected),
        "profile_overrides": deepcopy(profile_overrides or {}),
        "probe": [p.to_dict() for p in probes],
        "note": (
            "Bootstrapped from family rules only. No AZD project fixes or exceptions were inherited; "
            "project-specific knowledge must be learned through review."
        ),
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(profile, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return output_path, profile
