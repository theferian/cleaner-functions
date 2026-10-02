from __future__ import annotations

from cleaner_core.contracts import Finding
from cleaner_core.card_map import card_map_config, scan_all_card_maps
from openpyxl.utils import get_column_letter

TEST_META = {
    "id": "CARD_MAP",
    "title": "Card Data Boundary Map",
    "kind": "test",
    "profile_rule": "data_model.card_map",
    "snapshot_mode": "openpyxl",
}


def run(snapshot, profile, log):
    config = card_map_config(profile)
    log(
        "START",
        "CARD_MAP maps each DocumentRegion into a human-readable identifier area and data area. "
        "It does not validate dates/references yet.",
    )
    scans = scan_all_card_maps(snapshot, profile)
    findings: list[Finding] = []

    for sheet, scan in scans.items():
        start = get_column_letter(scan.data_start_col) if scan.data_start_col else "?"
        end = get_column_letter(scan.data_end_col) if scan.data_end_col else "?"
        with_data = sum(1 for x in scan.cards if x.slots)
        log(
            "SUMMARY",
            f"{sheet}: cards={len(scan.cards)}, data boundary={start}:{end}, "
            f"cards with mapped data slots={with_data}, issues={len(scan.issues)}.",
        )
        if scan.cards:
            example = scan.cards[0]
            log(
                "INFO",
                f"{sheet}: example card -> identifiers {example.identifier_range}; "
                f"data {example.data_range}; visible slots={len(example.slots)}.",
            )

        for issue in scan.issues:
            findings.append(Finding(
                test_id="CARD_MAP",
                status="FAIL",
                severity="HIGH",
                sheet=sheet,
                target=None,
                message=issue["message"],
                expected="Every derived document card has a usable data boundary",
                observed=issue,
                evidence=issue,
            ))

    pending_roles = [x for x in config["semantic_fields"] if x not in config["field_map"]]
    if pending_roles:
        log(
            "INFO",
            "Semantic field roles are intentionally not validated yet. Pending mappings: "
            + ", ".join(pending_roles),
        )

    log("DONE", f"CARD_MAP finished with {len(findings)} structural issue finding(s).")
    return findings
