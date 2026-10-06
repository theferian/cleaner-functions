from __future__ import annotations

from cleaner_core.contracts import Finding
from cleaner_core.field_schema import infer_all_slot_schemas
from openpyxl.utils import get_column_letter

TEST_META = {
    "id": "SLOT_SCHEMA",
    "title": "Transaction Slot Schema Discovery",
    "kind": "test",
    "profile_rule": "data_model.field_inference",
    "snapshot_mode": "openpyxl",
}


def run(snapshot, profile, log):
    log(
        "START",
        "SLOT_SCHEMA infers repeating transaction blocks and human-readable field roles from headers and cross-card majority patterns.",
    )
    schemas = infer_all_slot_schemas(snapshot, profile)
    findings: list[Finding] = []

    for sheet, schema in schemas.items():
        start = get_column_letter(schema.data_start_col) if schema.data_start_col else "?"
        end = get_column_letter(schema.data_end_col) if schema.data_end_col else "?"
        roles = ", ".join(
            f"+{key.replace(':', '/')}={role}"
            for key, role in sorted(
                schema.roles_by_slot.items(),
                key=lambda item: (int(item[0].split(":", 1)[0]), item[0].split(":", 1)[1]),
            )
        ) or "none"
        log(
            "SUMMARY",
            f"{sheet}: data={start}:{end}, block_width={schema.block_width}, "
            f"block_confidence={schema.block_confidence:.2f}, block_source={schema.block_source}, "
            f"roles=[{roles}].",
        )
        if schema.block_source == "SEMANTIC_ANCHOR_SPACING" and schema.block_evidence:
            log(
                "INFO",
                f"{sheet}: block width recovered from {schema.block_evidence.get('semantic_anchor')} "
                f"anchor columns {schema.block_evidence.get('columns')} with gaps "
                f"{schema.block_evidence.get('gaps')}.",
            )
        if schema.reference_rule:
            log(
                "INFO",
                f"{sheet}: inferred reference regex {schema.reference_rule.regex!r} "
                f"from {schema.reference_rule.matched_count}/{schema.reference_rule.sample_count} values "
                f"(coverage={schema.reference_rule.coverage:.2f}).",
            )
        else:
            log("INFO", f"{sheet}: no high-confidence reference regex inferred yet.")
        if schema.status_values:
            log(
                "INFO",
                f"{sheet}: inferred status vocabulary={list(schema.status_values)} "
                f"(coverage={schema.status_coverage:.2f}).",
            )

        mapped = [
            f"{p.letter}[{p.lane}]={p.inferred_role}({p.role_confidence:.2f})"
            for p in schema.slot_profiles if p.inferred_role
        ]
        suppressed = [
            f"{p.letter}[{p.lane}]={p.semantic_type}({p.semantic_confidence:.2f})"
            for p in schema.slot_profiles if p.semantic_type in {"MIXED", "UNKNOWN"}
        ]
        log("FIELDS", f"{sheet}: validation fields: {', '.join(mapped) if mapped else 'none'}.")
        if suppressed:
            log("SUPPRESSED", f"{sheet}: not validated: {', '.join(suppressed)}.")

        for issue in schema.issues:
            findings.append(Finding(
                test_id="SLOT_SCHEMA",
                status="POSSIBLE_EXCEPTION",
                severity="MEDIUM",
                sheet=sheet,
                target=f"{start}:{end}",
                message=issue["message"],
                expected="A repeatable, high-confidence transaction slot schema",
                observed=issue,
                evidence={**issue, "schema": schema.to_dict()},
            ))

    log("DONE", f"SLOT_SCHEMA finished with {len(findings)} schema review item(s).")
    return findings
