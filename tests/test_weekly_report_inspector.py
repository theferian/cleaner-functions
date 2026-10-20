from pathlib import Path
import importlib.util

import pytest
from openpyxl import Workbook

MODULE=Path(__file__).resolve().parents[1]/"tools"/"weekly_report"/"inspect_workbook.py"
spec=importlib.util.spec_from_file_location("wr_inventory_test_module",MODULE)
app=importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)


def test_jalali_candidate():
    assert app.JALALI.search("گزارش 1405/07/22")
    assert not app.JALALI.search("1450/99/99")


def test_inspect_preserves_source(tmp_path):
    path=tmp_path/"sample.xlsx"
    wb=Workbook();ws=wb.active;ws.title="Report";ws["A1"]="WBS";ws["B1"]="1405/07/22";ws["B2"]="=SUM(2,3)";ws.merge_cells("C3:D3")
    wb.create_sheet("Metadata");wb.save(path)
    before=path.read_bytes()
    report=app.inspect(path)
    assert path.read_bytes()==before
    assert report["sha256"]==app.hash_file(path)
    assert len(report["sheets"])==2
    assert report["sheets"][0]["formula_count"]==1
    assert report["sheets"][0]["jalali_date_candidates"][0]["cell"]=="B1"
    assert "C3:D3" in report["sheets"][0]["merged_ranges"]


def test_reject_wrong_format(tmp_path):
    path=tmp_path/"old.xls";path.write_text("x")
    with pytest.raises(ValueError):app.inspect(path)
