#!/usr/bin/env python3
"""Read-only discovery of an EPC Weekly Report XLSX/XLSM workbook.

Do not mutate source files or assume the template has been mapped.
"""
from __future__ import annotations
import argparse
import datetime as dt
import hashlib
import json
from pathlib import Path
import re

JALALI = re.compile(r"(?<!\d)(13\d{2}|14\d{2})[/.-](0?[1-9]|1[0-2])[/.-](0?[1-9]|[12]\d|3[01])(?!\d)")
LABELS = ("revision","rev","wbs","progress","planned","actual","date","status","engineering","procurement",
          "construction","برنامه","پیشرفت","تاریخ","وضعیت","نسخه","مهندسی","تدارکات","اجرا","گزارش")


def hash_file(path: Path):
    digest=hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_text(value, limit=150):
    if value is None:return ""
    if isinstance(value,(dt.datetime,dt.date)):
        return value.isoformat()
    v=str(value).strip().replace("\n"," ").replace("\r"," ")
    return v[:limit]+("..." if len(v)>limit else "")


def inspect(path: Path, sample_rows=25, sample_cols=28):
    from openpyxl import load_workbook
    if path.suffix.lower() not in (".xlsx",".xlsm"):
        raise ValueError("XLSX/XLSM required. For XLS, make a verified conversion copy in Excel first.")
    if not path.is_file():raise FileNotFoundError(path)
    wb=load_workbook(path,read_only=False,data_only=False,keep_links=True)
    result={"source_file":str(path.resolve()),"sha256":hash_file(path),"inspected_at_utc":dt.datetime.now(dt.timezone.utc).isoformat(),
            "read_only_inspection":True,"sheets":[],"named_ranges":[],"warnings":[]}
    try:
        for name in wb.defined_names:
            result["named_ranges"].append(str(name))
        for idx,ws in enumerate(wb.worksheets,1):
            rows=min(ws.max_row,sample_rows); cols=min(ws.max_column,sample_cols)
            preview=[]; jalali=[]; headers=[]; formulas=[]; links=[]
            total_nonempty=0; total_formulas=0; total_links=0
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:continue
                    total_nonempty+=1
                    if cell.data_type=="f":
                        total_formulas+=1
                        if len(formulas)<30:formulas.append({"cell":cell.coordinate,"expression":safe_text(cell.value)})
                    if cell.hyperlink:
                        total_links+=1
                        if len(links)<20:links.append({"cell":cell.coordinate,"target":safe_text(cell.hyperlink.target)})
                    s=safe_text(cell.value)
                    if len(jalali)<60 and JALALI.search(s):jalali.append({"cell":cell.coordinate,"value":s})
                    if len(headers)<60 and any(k in s.casefold() for k in LABELS) and len(s)<110:
                        headers.append({"cell":cell.coordinate,"text":s})
                    if cell.row<=rows and cell.column<=cols:
                        preview.append({"cell":cell.coordinate,"value":s,"type":cell.data_type})
            result["sheets"].append({"index":idx,"name":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,
                "nonempty_cells":total_nonempty,"formula_count":total_formulas,"hyperlink_count":total_links,
                "merged_ranges":[str(r) for r in list(ws.merged_cells.ranges)[:80]],
                "merged_ranges_total":len(ws.merged_cells.ranges),"headers":headers,
                "jalali_date_candidates":jalali,"formulas_sample":formulas,"external_links_sample":links,
                "first_region_nonempty_preview":preview})
    finally:wb.close()
    return result


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--file",required=True,help="Weekly Report XLSX/XLSM source")
    p.add_argument("--out",help="JSON output; defaults to working/weekly_report_discovery/<source>.inventory.json")
    p.add_argument("--sample-rows",type=int,default=25)
    p.add_argument("--sample-cols",type=int,default=28)
    a=p.parse_args()
    src=Path(a.file)
    out=Path(a.out) if a.out else Path("working/weekly_report_discovery")/(src.stem+".inventory.json")
    payload=inspect(src,a.sample_rows,a.sample_cols)
    out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(payload,indent=2,ensure_ascii=False),encoding="utf-8")
    print(f"[SOURCE] {src} SHA256={payload['sha256'][:12]}...")
    for s in payload["sheets"]:
        print(f"[SHEET] {s['index']}: {s['name']} rows={s['max_row']} cols={s['max_column']} merged={s['merged_ranges_total']} dates={len(s['jalali_date_candidates'])} formulas={s['formula_count']}")
    print(f"[READ_ONLY] {out}")

if __name__=="__main__":main()
