from pathlib import Path

from openpyxl import Workbook

from cleaner_core.card_map import scan_card_map
from cleaner_core.snapshot import WorkbookSnapshot
from tests.card_map import run as card_map_check
from cleaner_core.logging import RunLogger


def make_profile():
    return {
        "profile_name": "AZD",
        "structure": {"m_line": {"target_sheets": ["230KV SS"], "boundary_after_col": 3}},
        "data_model": {
            "family_fields": [
                {"key": "document_type", "default_locator": {"kind": "COLUMN", "column": 1}},
                {"key": "name", "default_locator": {"kind": "COLUMN", "column": 2}},
                {"key": "wbs", "default_locator": {"kind": "COLUMN", "column": 3}},
            ],
            "instances": {"230KV SS": {"field_map": {
                "document_type": {"kind": "COLUMN", "column": 1},
                "name": {"kind": "COLUMN", "column": 2},
                "wbs": {"kind": "COLUMN", "column": 3},
            }}},
            "document_boundaries": {
                "first_data_row": 3,
                "signal_fields": ["document_type", "name", "wbs"],
                "boundary_fields": ["name", "wbs"],
                "separator_fields": ["name", "wbs"],
            },
            "card_map": {
                "enabled": True,
                "data_start_col": 4,
                "data_end_strategy": "MERGE_EXTENT",
                "semantic_fields": ["reference", "send_date", "receive_date", "status"],
                "field_map": {},
            },
        },
    }


def build_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for row, name, wbs in [(3, "Doc A", "1.1"), (5, "Doc B", "1.2")]:
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = name
        ws[f"C{row}"] = wbs
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
        ws[f"D{row}"] = "REF"
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        ws[f"G{row}"] = "1405/01/01"
        ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=8)
    wb.save(path)


def test_card_map_exposes_identifier_and_data_boundaries(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)
    scan = scan_card_map(snap, make_profile(), "230KV SS")
    assert len(scan.cards) == 2
    assert scan.data_start_col == 4
    assert scan.data_end_col == 8
    assert scan.cards[0].identifier_range == "A3:C4"
    assert scan.cards[0].data_range == "D3:H4"
    assert {x.address for x in scan.cards[0].slots} >= {"D3:F3", "G3:H3"}
    assert scan.issues == []
    snap.close()


def test_card_map_stage_does_not_force_semantic_validation_yet(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)
    findings = card_map_check(snap, make_profile(), RunLogger(stream=False))
    assert findings == []
    snap.close()
