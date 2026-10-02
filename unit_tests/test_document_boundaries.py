from pathlib import Path

from openpyxl import Workbook

from cleaner_core.document_regions import scan_document_regions
from cleaner_core.snapshot import WorkbookSnapshot


def make_profile():
    return {
        "structure": {"m_line": {"target_sheets": ["230KV SS"]}},
        "data_model": {
            "family_fields": [
                {
                    "key": "document_type",
                    "default_locator": {"kind": "COLUMN", "column": 1},
                },
                {
                    "key": "name",
                    "default_locator": {"kind": "COLUMN", "column": 2},
                },
                {
                    "key": "wbs",
                    "default_locator": {"kind": "COLUMN", "column": 3},
                },
            ],
            "instances": {
                "230KV SS": {
                    "field_map": {
                        "document_type": {"kind": "COLUMN", "column": 1},
                        "name": {"kind": "COLUMN", "column": 2},
                        "wbs": {"kind": "COLUMN", "column": 3},
                    }
                }
            },
            "document_boundaries": {
                "first_data_row": 3,
                "signal_fields": ["document_type", "name", "wbs"],
                "boundary_fields": ["name", "wbs"],
                "separator_fields": ["name", "wbs"],
            },
        },
    }


def build_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"

    # A two-row document: Name and WBS own their vertical merge anchors.
    ws["A3"] = "DOC"
    ws["B3"] = "Document A"
    ws["C3"] = "1.2.3"
    ws.merge_cells("B3:B4")
    ws.merge_cells("C3:C4")

    # A one-row document.
    ws["A5"] = "DOC"
    ws["B5"] = "Document B"
    ws["C5"] = "1.2.4"

    # Background/category separator crossing the identifier columns. It must
    # not become a document region.
    ws["A6"] = "CATEGORY"
    ws.merge_cells("A6:F7")

    # Empty Name/WBS is still a document candidate because document_type is a
    # mapped signal field. This preserves the ability to detect missing IDs.
    ws["A8"] = "DOC"

    wb.save(path)


def test_document_region_spans_and_separator_exclusion(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)

    scan = scan_document_regions(snap, make_profile(), "230KV SS")

    assert [(r.start_row, r.end_row) for r in scan.regions] == [
        (3, 4),
        (5, 5),
        (8, 8),
    ]
    assert scan.regions[0].field("name").merge == "B3:B4"
    assert "A6:F7" in scan.excluded_separator_ranges
    assert scan.ambiguities == []

    snap.close()


def test_external_identifier_merge_is_ambiguous(tmp_path):
    path = tmp_path / "ambiguous.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    ws["A3"] = "CATEGORY"
    ws.merge_cells("A3:B3")  # covers mapped Name from an external anchor
    ws["C3"] = "1.2.3"      # WBS still makes the row a candidate
    wb.save(path)

    snap = WorkbookSnapshot(path)
    scan = scan_document_regions(snap, make_profile(), "230KV SS")

    assert len(scan.regions) == 1
    assert len(scan.ambiguities) == 1
    reasons = scan.ambiguities[0]["reasons"]
    assert any(x["kind"] == "EXTERNAL_IDENTIFIER_MERGE" for x in reasons)

    snap.close()
