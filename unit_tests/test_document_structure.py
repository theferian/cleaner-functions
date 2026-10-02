from pathlib import Path

from openpyxl import Workbook

from cleaner_core.document_structure import scan_document_structure, structure_pattern_for_region
from cleaner_core.knowledge import promote_decisions_to_profile
from cleaner_core.document_regions import DocumentRegion, FieldAnchor
from cleaner_core.ranges import parse_range
from cleaner_core.logging import RunLogger
from cleaner_core.snapshot import WorkbookSnapshot
from tests.document_structure import run as document_structure_check


def make_profile():
    return {
        "profile_name": "AZD",
        "structure": {
            "m_line": {
                "target_sheets": ["230KV SS"],
                "boundary_after_col": 3,
            }
        },
        "data_model": {
            "family_fields": [
                {"key": "document_type", "default_locator": {"kind": "COLUMN", "column": 1}},
                {"key": "name", "default_locator": {"kind": "COLUMN", "column": 2}},
                {"key": "wbs", "default_locator": {"kind": "COLUMN", "column": 3}},
            ],
            "instances": {
                "230KV SS": {
                    "field_map": {
                        "document_type": {"kind": "COLUMN", "column": 1},
                        "name": {"kind": "COLUMN", "column": 2},
                        "wbs": {"kind": "COLUMN", "column": 3},
                    }
                }
            },
            "document_boundaries": {
                "first_data_row": 3,
                "signal_fields": ["document_type", "name", "wbs"],
                "boundary_fields": ["name", "wbs"],
                "separator_fields": ["name", "wbs"],
            },
            "document_structure": {
                "enabled": True,
                "method": "NORMALIZED_MERGE_SIGNATURE",
                "transaction_start_col": 4,
                "unknown_pattern_status": "POSSIBLE_EXCEPTION",
                "known_patterns": [],
            },
        },
        "known_exceptions": [],
        "known_fixes": [],
    }


def build_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"

    # Card A: rows 3:4.
    ws["A3"] = "DOC"
    ws["B3"] = "Document A"
    ws["C3"] = "1.1"
    ws.merge_cells("B3:B4")
    ws.merge_cells("C3:C4")
    ws.merge_cells("D3:F3")
    ws.merge_cells("D4:E4")

    # Card B: same geometry, shifted to rows 5:6.
    ws["A5"] = "DOC"
    ws["B5"] = "Document B"
    ws["C5"] = "1.2"
    ws.merge_cells("B5:B6")
    ws.merge_cells("C5:C6")
    ws.merge_cells("D5:F5")
    ws.merge_cells("D6:E6")

    # Card C: different transaction merge geometry.
    ws["A7"] = "DOC"
    ws["B7"] = "Document C"
    ws["C7"] = "1.3"
    ws.merge_cells("B7:B8")
    ws.merge_cells("C7:C8")
    ws.merge_cells("D7:G7")
    ws.merge_cells("D8:E8")

    # Legitimate higher-level/context merges that span several document cards.
    # Column A is contextual (Document Type), while H is in the transaction side.
    # Neither defines a DocumentRegion boundary, so these must not become hazards.
    ws.merge_cells("A3:A8")
    ws.merge_cells("H3:H8")

    wb.save(path)


def test_structure_signature_is_position_independent(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)
    scan = scan_document_structure(snap, make_profile(), "230KV SS")

    assert len(scan.regions) == 3
    ids = [x.pattern.pattern_id for x in scan.regions]
    assert ids[0] == ids[1]
    assert ids[2] != ids[0]
    assert scan.patterns[ids[0]]["count"] == 2
    assert scan.hazards == []
    snap.close()


def test_unknown_patterns_are_reviewed_once_per_geometry(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)
    log = RunLogger(stream=False)

    findings = document_structure_check(snap, make_profile(), log)
    unknown = [x for x in findings if x.status == "POSSIBLE_EXCEPTION"]
    assert len(unknown) == 2
    assert sorted(x.evidence["pattern_count"] for x in unknown) == [1, 2]
    snap.close()


def test_reviewed_structure_pattern_can_be_promoted_and_recognized(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    profile = make_profile()
    snap = WorkbookSnapshot(path)
    log = RunLogger(stream=False)
    findings = document_structure_check(snap, profile, log)

    repeated = next(x for x in findings if x.evidence["pattern_count"] == 2)
    payload = {
        "run": {"run_id": "week_1", "script_id": "DOCUMENT_STRUCTURE"},
        "decisions": [
            {
                "review_id": "s1",
                "source": {
                    "test_id": "DOCUMENT_STRUCTURE",
                    "status": repeated.status,
                    "severity": repeated.severity,
                },
                "location": {"sheet": repeated.sheet, "target": repeated.target},
                "evidence": repeated.evidence,
                "decision": {
                    "action": "ACCEPT_EXCEPTION",
                    "parameters": {
                        "test_id": "DOCUMENT_STRUCTURE",
                        "sheet": repeated.sheet,
                        "target": repeated.target,
                    },
                    "reason": "Reviewed as valid AZD card geometry",
                },
            }
        ],
    }
    updated, report = promote_decisions_to_profile(profile, payload)
    assert report["promoted_structure_pattern_count"] == 1
    assert report["promoted_exception_count"] == 0

    findings2 = document_structure_check(snap, updated, RunLogger(stream=False))
    unknown2 = [x for x in findings2 if x.status == "POSSIBLE_EXCEPTION"]
    assert len(unknown2) == 1
    assert unknown2[0].evidence["pattern_count"] == 1
    snap.close()



def test_owned_boundary_merge_crossing_region_is_a_hazard():
    region = DocumentRegion(
        sheet="230KV SS",
        record_index=1,
        anchor_row=3,
        start_row=3,
        end_row=4,
        fields={
            "name": FieldAnchor(
                field_key="name", column=2, address="B3", value="Document A",
                own_anchor=True, continuation=False, covered_by_external_merge=False,
                merge="B3:B4", span_start_row=3, span_end_row=4,
            ),
            "wbs": FieldAnchor(
                field_key="wbs", column=3, address="C3", value="1.1",
                own_anchor=True, continuation=False, covered_by_external_merge=False,
                merge="C3:C4", span_start_row=3, span_end_row=4,
            ),
        },
        signal_fields=("name", "wbs"),
        boundary_fields=("name", "wbs"),
    )

    _, crossing = structure_pattern_for_region(
        region=region,
        merged_ranges=[parse_range("B3:B6")],
        transaction_start_col=4,
    )

    assert crossing == ("B3:B6",)
