from pathlib import Path

from openpyxl import Workbook

from cleaner_core.logging import RunLogger
from cleaner_core.snapshot import WorkbookSnapshot
from tests.field_validation import run as field_validation_check


def make_profile():
    return {
        "profile_name": "AZD",
        "structure": {"m_line": {"target_sheets": ["230KV SS"], "boundary_after_col": 3}},
        "data_model": {
            "family_fields": [
                {"key": "document_type", "default_locator": {"kind": "COLUMN", "column": 1}},
                {"key": "name", "default_locator": {"kind": "COLUMN", "column": 2}},
                {"key": "wbs", "default_locator": {"kind": "COLUMN", "column": 3}},
            ],
            "instances": {"230KV SS": {"field_map": {
                "document_type": {"kind": "COLUMN", "column": 1},
                "name": {"kind": "COLUMN", "column": 2},
                "wbs": {"kind": "COLUMN", "column": 3},
            }}},
            "document_boundaries": {
                "first_data_row": 3,
                "signal_fields": ["document_type", "name", "wbs"],
                "boundary_fields": ["name", "wbs"],
                "separator_fields": ["name", "wbs"],
            },
            "card_map": {"data_start_col": 4},
            "field_inference": {
                "header_rows": [1, 2],
                "max_block_width": 6,
                "min_block_confidence": 0.65,
                "status_min_occurrences": 2,
                "status_min_coverage": 0.70,
            },
        },
    }


def build_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for col, header in [(4, "Send Date"), (5, "Receive Date"), (6, "Status"),
                        (7, "Send Date"), (8, "Receive Date"), (9, "Status")]:
        ws.cell(2, col).value = header
    for i, row in enumerate([3, 5, 7, 9]):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
        ws[f"D{row}"] = "1404/01/01"
        ws[f"E{row}"] = "1404/01/02"
        ws[f"F{row}"] = "CM"
        ws[f"G{row}"] = "1404/02/01"
        ws[f"H{row}"] = "1404/02/02"
        ws[f"I{row}"] = "CM"
        ws.merge_cells(start_row=row, start_column=6, end_row=row + 1, end_column=6)
        ws.merge_cells(start_row=row, start_column=9, end_row=row + 1, end_column=9)
    # One deterministic normalization and one genuinely invalid date.
    ws["D5"] = "۱۴۰۴-۱-۳"
    ws["H7"] = "1404/13/01"
    # One rare status should enter review.
    ws["I9"] = "ZZ"
    wb.save(path)


def test_field_validation_only_returns_failed_or_noncanonical_values(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, make_profile(), RunLogger(stream=False))
    targets = {x.target: x for x in findings}
    assert set(targets) == {"D5", "E5", "H7", "I9:I10"}
    assert targets["D5"].suggested_fix["parameters"]["value"] == "1404/01/03"
    assert targets["E5"].evidence["audit_kind"] == "DATE_CHRONOLOGY"
    assert targets["E5"].suggested_fix is None
    assert targets["H7"].suggested_fix is None
    assert targets["I9:I10"].suggested_fix is None
    assert all(x.status == "FAIL" for x in findings)
    snap.close()


def test_multiline_card_dispatches_each_physical_subslot_by_value_shape(tmp_path):
    path = tmp_path / "stacked_card.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for col, header in [(4, "Send Date"), (5, "Receive Date"), (6, "Status"),
                        (7, "Send Date"), (8, "Receive Date"), (9, "Status")]:
        ws.cell(2, col).value = header

    # Several ordinary two-row cards establish the repeating 3-column rhythm.
    for i, row in enumerate([3, 5, 7, 9]):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
        for start_col, prefix in [(4, "ES"), (7, "ES")]:
            ws.cell(row, start_col).value = f"1404/03/{i+1:02d}"
            ws.cell(row + 1, start_col).value = f"1404/{prefix}/S/{1600+i}"
            ws.cell(row, start_col + 1).value = f"1404/04/{i+1:02d}"
            ws.cell(row + 1, start_col + 1).value = f"1404/{prefix}/R/{1700+i}"
            ws.cell(row, start_col + 2).value = "CM"
            ws.merge_cells(start_row=row, start_column=start_col + 2, end_row=row + 1, end_column=start_col + 2)

    # Four-row special card: same column stacks date -> reference -> date -> reference.
    row = 11
    ws[f"A{row}"] = "DOC"
    ws[f"B{row}"] = "Special"
    ws[f"C{row}"] = "1.99"
    ws.merge_cells(start_row=row, start_column=2, end_row=row + 3, end_column=2)
    ws.merge_cells(start_row=row, start_column=3, end_row=row + 3, end_column=3)
    ws["G11"] = "1404/03/05"
    ws["G12"] = "1404/ES/S/1601"
    ws["G13"] = "1404/06/30"
    ws["G14"] = "1404/ES/S/1602"
    ws["H11"] = "عطف به بند شماره 1 صورتجلسه مورخ 04/06/03"
    ws.merge_cells("H11:H12")
    ws["H13"] = "1404/07/01"
    ws["I11"] = "CM"
    ws.merge_cells("I11:I14")
    wb.save(path)

    profile = make_profile()
    profile["data_model"]["field_inference"].update({
        "reference_min_samples": 3,
        "reference_min_coverage": 0.60,
        "semantic_min_confidence": 0.80,
    })
    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, profile, RunLogger(stream=False))
    by_target = {f.target: f for f in findings}

    # These two were the real regression: reference must not be tested as a date,
    # and a later stacked date must not be tested as a reference.
    assert "G12" not in by_target
    assert "G13" not in by_target

    # The meeting-note explanation is preserved; only its embedded short date is normalized.
    assert "H11:H12" in by_target
    assert by_target["H11:H12"].evidence["field_role"] == "explanation"
    assert by_target["H11:H12"].suggested_fix["parameters"]["value"] == (
        "عطف به بند شماره 1 صورتجلسه مورخ 1404/06/03"
    )
    snap.close()


def test_global_date_audit_catches_dates_outside_inferred_date_fields(tmp_path):
    path = tmp_path / "global_date_audit.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for i, row in enumerate([3, 5, 7]):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
    # These are deliberately narrative/unmapped slots, not inferred date fields.
    ws["D3"] = "note 1404/13/01"
    ws["D5"] = "note 1404-2-3"
    ws["D7"] = "note 1404/02/03"
    wb.save(path)

    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, make_profile(), RunLogger(stream=False))
    by_target = {f.target: f for f in findings}
    assert "D3" in by_target
    assert by_target["D3"].evidence["audit_kind"] == "GLOBAL_DATE_TOKEN_SCAN"
    assert by_target["D3"].suggested_fix is None
    assert "D5" in by_target
    assert by_target["D5"].suggested_fix["parameters"]["value"] == "note 1404/02/03"
    assert "D7" not in by_target
    snap.close()


def test_receive_date_requires_corresponding_status(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "receive_status.xlsx"
    build_book(path)
    wb = load_workbook(path)
    ws = wb["230KV SS"]
    # F3:F4 is the status slot corresponding to E3 receive date.
    ws["F3"] = None
    wb.save(path)
    wb.close()

    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, make_profile(), RunLogger(stream=False))
    relations = [f for f in findings if f.evidence.get("audit_kind") == "RECEIVE_STATUS_RELATION"]
    assert any(f.target == "E3" for f in relations)
    snap.close()



def test_global_date_audit_discovers_malformed_components(tmp_path):
    path = tmp_path / "broad_global_date_audit.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for i, row in enumerate([3, 5, 7]):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
    # Month has three digits: discovery must still see it, then calendar validation rejects it.
    ws["D3"] = "note 1404/013/01"
    # Three-digit year is a likely typo and should be surfaced, not silently missed.
    ws["D5"] = "note 404/02/03"
    # Five-digit 14xxx year is also a likely typing error (e.g. 14040/09/23).
    ws["D7"] = "note 14040/09/23"
    wb.save(path)

    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, make_profile(), RunLogger(stream=False))
    global_findings = [f for f in findings if f.evidence.get("audit_kind") == "GLOBAL_DATE_TOKEN_SCAN"]
    targets = {f.target for f in global_findings}
    assert "D3" in targets
    assert "D5" in targets
    assert "D7" in targets
    snap.close()


def test_date_chronology_flags_receive_before_send(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "chronology.xlsx"
    build_book(path)
    wb = load_workbook(path)
    ws = wb["230KV SS"]
    ws["D3"] = "1404/05/10"  # send
    ws["E3"] = "1404/05/09"  # receive before send
    wb.save(path)
    wb.close()

    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, make_profile(), RunLogger(stream=False))
    chronology = [f for f in findings if f.evidence.get("audit_kind") == "DATE_CHRONOLOGY"]
    assert any(f.target == "E3" for f in chronology)
    assert chronology[0].suggested_fix is None
    snap.close()


def test_known_status_value_overrides_irregular_card_geometry(tmp_path):
    """Known status vocabulary wins when an irregular card puts status in a receive slot."""
    path = tmp_path / "status_geometry_override.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for col, header in [(4, "Send Date"), (5, "Receive Date"), (6, "Status"),
                        (7, "Send Date"), (8, "Receive Date"), (9, "Status")]:
        ws.cell(2, col).value = header

    # Normal cards establish the transaction rhythm and A.N as a real status.
    for i, row in enumerate([3, 5, 7, 9]):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
        ws[f"D{row}"] = "1404/01/01"
        ws[f"E{row}"] = "1404/01/02"
        ws[f"F{row}"] = "A.N" if i < 2 else "CM"
        ws.merge_cells(start_row=row, start_column=6, end_row=row + 1, end_column=6)
        ws[f"G{row}"] = "1404/02/01"
        ws[f"H{row}"] = "1404/02/02"
        ws[f"I{row}"] = "CM"
        ws.merge_cells(start_row=row, start_column=9, end_row=row + 1, end_column=9)

    # Exceptional geometry: A.N is physically in a RECEIVE column. It must be
    # read as status, not sent through reference/date validation.
    ws["E9"] = "A.N"
    wb.save(path)

    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, make_profile(), RunLogger(stream=False))
    assert not any(
        f.target == "E9" and f.evidence.get("audit_kind") == "REFERENCE_REGEX"
        for f in findings
    )
    assert not any(
        f.target == "E9" and f.evidence.get("audit_kind") == "SEMANTIC_DATE"
        for f in findings
    )
    snap.close()


def test_reusable_status_library_accepts_rare_rj(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "rare_rj.xlsx"
    build_book(path)
    wb = load_workbook(path)
    ws = wb["230KV SS"]
    # Replace the synthetic unknown ZZ with the user-confirmed catalogued status RJ.
    ws["I9"] = "RJ"
    wb.save(path)
    wb.close()

    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, make_profile(), RunLogger(stream=False))
    by_target = {f.target: f for f in findings}
    assert "I9:I10" not in by_target
    snap.close()

def test_reusable_status_library_accepts_fsh_alias_and_canonical_status(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "fsh_statuses.xlsx"
    build_book(path)
    wb = load_workbook(path)
    ws = wb["230KV SS"]
    ws["I9"] = "C.M"
    ws["F9"] = "N.A"
    wb.save(path)
    wb.close()

    snap = WorkbookSnapshot(path)
    findings = field_validation_check(snap, make_profile(), RunLogger(stream=False))
    by_target = {f.target: f for f in findings}
    assert "I9:I10" not in by_target
    assert "F9:F10" not in by_target
    snap.close()
