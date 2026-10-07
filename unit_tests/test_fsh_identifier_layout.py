from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from cleaner_core.logging import RunLogger
from cleaner_core.snapshot import WorkbookSnapshot
from functions.cleaner.bootstrap import build_general_document_profile
from tests.identifier_check import run as identifier_check


def _fsh_overrides(root: Path) -> dict:
    cfg = json.loads((root / "configs" / "project_control_sources.json").read_text(encoding="utf-8"))
    return cfg["projects"]["FSH"]["cleaner_profiles"]["GENERAL_DOCUMENT"]


def _book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "FOULAD SHAHRIYAR"
    ws["A1"] = "Category"
    ws["B1"] = "Type"
    ws["C1"] = "Name"
    ws["D1"] = "Code"
    ws["E1"] = "WBS"
    ws["F1"] = "Send Rev Z"
    ws["G1"] = "Receive"
    ws["H1"] = "Status"

    rows = [
        ("CAT", "TYPE", "Doc A", "ENG.10", "1.1", "1405/01/01"),
        ("CAT", "TYPE", "Doc B", None, "1.2", "1405/01/02"),
        ("CAT", "TYPE", "Doc C", "-", "1.3", "1405/01/03"),
        ("CAT", "TYPE", "Doc D", "EL.5", "2.1", "1405/01/04"),
        ("CAT", "TYPE", "Doc E", None, "2.2", "1405/01/05"),
    ]
    for row, values in enumerate(rows, start=3):
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col).value = value
    wb.save(path)


def test_fsh_profile_override_sets_five_identifier_columns(tmp_path):
    root = Path(__file__).resolve().parents[1]
    workbook = tmp_path / "fsh.xlsx"
    _book(workbook)
    out = tmp_path / "FSH.profile.json"
    _, profile = build_general_document_profile(
        project_id="FSH",
        workbook_path=workbook,
        base_profile_path=root / "profiles" / "families" / "GENERAL_DOCUMENT.base.json",
        output_path=out,
        profile_overrides=_fsh_overrides(root),
    )
    assert profile["structure"]["m_line"]["boundary_after_col"] == 5
    assert profile["structure"]["m_line"]["boundary_label"] == "E|F"
    fmap = profile["data_model"]["instances"]["FOULAD SHAHRIYAR"]["field_map"]
    assert {k: v["column"] for k, v in fmap.items()} == {
        "category": 1,
        "type": 2,
        "name": 3,
        "code": 4,
        "wbs": 5,
    }
    assert profile["data_model"]["card_map"]["data_start_col"] == 6


def test_fsh_empty_code_uses_wbs_style_independent_sequence(tmp_path):
    root = Path(__file__).resolve().parents[1]
    workbook = tmp_path / "fsh.xlsx"
    _book(workbook)
    out = tmp_path / "FSH.profile.json"
    _, profile = build_general_document_profile(
        project_id="FSH",
        workbook_path=workbook,
        base_profile_path=root / "profiles" / "families" / "GENERAL_DOCUMENT.base.json",
        output_path=out,
        profile_overrides=_fsh_overrides(root),
    )
    snap = WorkbookSnapshot(workbook)
    try:
        findings = identifier_check(snap, profile, RunLogger(stream=False))
    finally:
        snap.close()

    code4 = next(x for x in findings if x.target == "D4" and x.evidence.get("field_key") == "code")
    code5 = next(x for x in findings if x.target == "D5" and x.evidence.get("field_key") == "code")
    code7 = next(x for x in findings if x.target == "D7" and x.evidence.get("field_key") == "code")
    assert code4.suggested_fix["parameters"]["value"] == "ENG.10.1000"
    assert code5.suggested_fix["parameters"]["value"] == "ENG.10.1001"
    assert code7.suggested_fix["parameters"]["value"] == "EL.5.1000"
    assert not any(x.evidence.get("field_key") == "wbs" and x.status == "FAIL" for x in findings)
