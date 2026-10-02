from pathlib import Path

from openpyxl import Workbook

from cleaner_core.snapshot import WorkbookSnapshot
from cleaner_core.logging import RunLogger
from tests.identifier_check import run as identifier_check


def make_profile():
    return {
        "schema_version": "0.4",
        "structure": {
            "m_line": {
                "target_sheets": ["230KV SS"]
            }
        },
        "data_model": {
            "family_fields": [
                {
                    "key": "document_type",
                    "label": "Document Type",
                    "data_type": "TEXT",
                    "required": False,
                    "default_locator": {"kind": "COLUMN", "column": 1},
                    "validation": {"kind": "NONE"},
                    "cleaning": {},
                },
                {
                    "key": "name",
                    "label": "Name",
                    "data_type": "TEXT",
                    "required": True,
                    "default_locator": {"kind": "COLUMN", "column": 2},
                    "validation": {"kind": "NON_EMPTY"},
                    "cleaning": {
                        "template": "ERR_EMPT_NAME_{NUMBER}",
                        "number_strategy": "ROW_NUMBER",
                    },
                },
                {
                    "key": "wbs",
                    "label": "WBS",
                    "data_type": "WBS",
                    "required": True,
                    "default_locator": {"kind": "COLUMN", "column": 3},
                    "validation": {
                        "kind": "REGEX",
                        "pattern": r"^\d+(?:\.\d+)*$",
                    },
                    "cleaning": {
                        "sequence_start": 1000,
                        "separator": ".",
                        "fallback_root_start": 1000,
                    },
                },
            ],
            "instances": {
                "230KV SS": {
                    "field_map": {
                        "document_type": {"kind": "COLUMN", "column": 1},
                        "name": {"kind": "COLUMN", "column": 2},
                        "wbs": {"kind": "COLUMN", "column": 3},
                    }
                }
            },
            "record_detection": {
                "first_data_row": 3,
                "signal_fields": ["document_type", "name", "wbs"],
            },
        },
    }


def build_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"

    # valid record
    ws["A3"] = "DOC"
    ws["B3"] = "Document A"
    ws["C3"] = "1.2.3"

    # empty name + empty WBS -> both suggestions
    ws["A4"] = "DOC"

    # invalid WBS should use same nearest valid parent and increment
    ws["A5"] = "DOC"
    ws["B5"] = "Document B"
    ws["C5"] = "-"

    # new valid parent
    ws["A6"] = "DOC"
    ws["B6"] = "Document C"
    ws["C6"] = "2.5"

    # invalid under new parent resets to .1000
    ws["A7"] = "DOC"
    ws["B7"] = "Document D"

    # separator/background merge should be ignored
    ws["A8"] = "SEPARATOR"
    ws.merge_cells("A8:F9")

    wb.save(path)


def test_identifier_suggestions(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)
    log = RunLogger(stream=False)

    findings = identifier_check(snap, make_profile(), log)

    name4 = next(
        x for x in findings
        if x.target == "B4" and x.evidence.get("field_key") == "name"
    )
    assert name4.suggested_fix["parameters"]["value"] == "ERR_EMPT_NAME_4"

    wbs4 = next(
        x for x in findings
        if x.target == "C4" and x.evidence.get("field_key") == "wbs"
    )
    assert wbs4.suggested_fix["parameters"]["value"] == "1.2.3.1000"

    wbs5 = next(
        x for x in findings
        if x.target == "C5" and x.evidence.get("field_key") == "wbs"
    )
    assert wbs5.suggested_fix["parameters"]["value"] == "1.2.3.1001"

    wbs7 = next(
        x for x in findings
        if x.target == "C7" and x.evidence.get("field_key") == "wbs"
    )
    assert wbs7.suggested_fix["parameters"]["value"] == "2.5.1000"

    assert not any(x.target in {"B8", "C8", "B9", "C9"} for x in findings)

    snap.close()
