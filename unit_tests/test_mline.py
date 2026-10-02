from pathlib import Path
import json
from openpyxl import Workbook

from cleaner_core.snapshot import WorkbookSnapshot
from cleaner_core.logging import RunLogger
from tests.m_line_check import run as m_line_check
from discoveries.discover_m_line import run as discover_m_line


def build_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"

    ws["A1"] = "AZD TITLE"
    ws.merge_cells("A1:C1")

    # Known exception geometry at rows 15-16.
    ws["C15"] = "Category/Type"
    ws.merge_cells("C15:D16")

    # Similar candidate geometry later.
    ws["C50"] = "Category/Type"
    ws.merge_cells("C50:D51")

    # Unexpected crossing.
    ws["B70"] = "Unexpected"
    ws.merge_cells("B70:E70")

    wb.create_sheet("63KV SS")
    wb.create_sheet("LINE")
    wb.save(path)


def profile():
    return {
        "schema_version":"0.3",
        "profile_name":"AZD",
        "structure":{
            "m_line":{
                "enabled":True,
                "boundary_after_col":3,
                "target_sheets":["230KV SS"],
                "candidate_exception_detection":{"enabled":True},
                "discovery":{"title_window":{"row_start":1,"row_end":10,"col_start":1,"col_end":6}}
            }
        },
        "known_exceptions":[
            {
                "id":"KNOWN",
                "enabled":True,
                "match":{"test_id":"M_LINE","sheet":"*","rows":[15,16]}
            }
        ]
    }


def test_discovery_and_check(tmp_path):
    path = tmp_path / "azd.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)
    log = RunLogger(stream=False)

    candidates = discover_m_line(snap, profile(), log)
    assert any(x.target == "A1:C1" for x in candidates)

    findings = m_line_check(snap, profile(), log)
    statuses = [x.status for x in findings]
    assert "KNOWN_EXCEPTION" in statuses
    assert "POSSIBLE_EXCEPTION" in statuses
    assert "FAIL" in statuses
    snap.close()
