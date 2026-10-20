from openpyxl import Workbook

from cleaner_core.snapshot import WorkbookSnapshot
from cleaner_core.logging import RunLogger
from review.queue import build_review_queue
from tests.m_line_check import run


def test_general_document_recognizes_full_width_group_separators(tmp_path):
    path = tmp_path / "YSK.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Yasookand"
    # Valid document above and below each divider.
    ws["B3"], ws["C3"] = "Document 1", "1.2.1"
    ws["B19"], ws["C19"] = "Document 2", "1.2.2"
    ws["B243"], ws["C243"] = "Document 3", "1.2.3"
    ws["B459"], ws["C459"] = "Document 4", "1.2.4"
    for start, label in ((17, "طراح"), (241, None), (457, None)):
        ws.merge_cells(start_row=start, start_column=1, end_row=start+1, end_column=30)
        ws.cell(start, 1).value = label
    # This is NOT a separator and must still fail M_LINE.
    ws["B470"] = "Merged actual content"
    ws.merge_cells("B470:E470")
    wb.save(path)

    snap = WorkbookSnapshot(path)
    profile = {"family": "GENERAL_DOCUMENT", "profile_name": "YSK", "structure": {"m_line": {
        "boundary_after_col": 3, "target_sheets": ["Yasookand"]
    }}}
    findings = run(snap, profile, RunLogger(stream=False))
    assert {f.target for f in findings if f.status == "PASS"} == {
        "A17:AD18", "A241:AD242", "A457:AD458"
    }
    assert {f.target for f in findings if f.status == "FAIL"} == {"B470:E470"}
    for f in findings:
        if f.status == "PASS":
            assert f.evidence["rule"] == "GENERAL_DOCUMENT_FULL_WIDTH_TWO_ROW_SEPARATOR"
    queue = build_review_queue(report={"findings": [f.to_dict() for f in findings]},
                               profile=profile, package_manifest={}, workbook={})
    assert queue["summary"]["total_items"] == 1
    snap.close()


def test_family_rule_does_not_swallow_nonfamily_or_date_bearing_or_opt_out(tmp_path):
    path = tmp_path / "ambiguous.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Docs"
    ws.merge_cells("A17:AD18")
    ws["A17"] = "1404/08/11: received letter"
    ws.merge_cells("A21:AD22")
    ws["A21"] = "Category"
    ws.merge_cells("A25:Q26")  # Partial-width merges must be reviewed.
    wb.save(path)
    snap = WorkbookSnapshot(path)
    base = {"family": "GENERAL_DOCUMENT", "structure": {"m_line": {
        "boundary_after_col": 3, "target_sheets": ["Docs"]
    }}}
    findings = run(snap, base, RunLogger(stream=False))
    assert {f.target for f in findings if f.status == "PASS"} == {"A21:AD22"}
    assert {f.target for f in findings if f.status == "FAIL"} == {"A17:AD18", "A25:Q26"}

    base["structure"]["m_line"]["family_horizontal_separators"] = {"enabled": False}
    findings = run(snap, base, RunLogger(stream=False))
    assert all(f.status == "FAIL" for f in findings)

    base["structure"]["m_line"]["family_horizontal_separators"] = {"enabled": True}
    base["family"] = "OTHER_FAMILY"
    findings = run(snap, base, RunLogger(stream=False))
    assert all(f.status == "FAIL" for f in findings)
    snap.close()
