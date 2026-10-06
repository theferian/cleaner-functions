from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from functions.cleaner.bootstrap import (
    build_general_document_profile,
    inspect_general_document_workbook,
)


def _workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "SUBSTATION"
    ws["A1"] = "Document Type"
    ws["B1"] = "Name"
    ws["C1"] = "WBS"
    ws["D1"] = "Send Rev Z"
    ws["E1"] = "Receive"
    ws["F1"] = "Status"
    ws["B3"] = "DOC A"
    ws["C3"] = "1.1"
    ws["D3"] = "1405/01/01"
    ws["B4"] = "DOC B"
    ws["C4"] = "1.2"
    ws["D4"] = "1405/01/02"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "hello"
    wb.save(path)


def test_bootstrap_discovers_data_sheets_without_azd_knowledge(tmp_path):
    workbook = tmp_path / "jnd.xlsx"
    _workbook(workbook)
    probes = inspect_general_document_workbook(workbook)
    selected = [x.sheet for x in probes if x.selected]
    assert selected == ["SUBSTATION"]


def test_bootstrap_creates_project_specific_profile(tmp_path):
    workbook = tmp_path / "fsh.xlsx"
    _workbook(workbook)
    base = {
        "schema_version": "0.4",
        "family": "GENERAL_DOCUMENT",
        "profile_name": "__PROJECT__",
        "workbook": {"expected_sheets": []},
        "structure": {"m_line": {"target_sheets": [], "boundary_after_col": 3}},
        "known_exceptions": [],
        "known_fixes": [],
        "data_model": {"family_fields": [], "instances": {}},
    }
    base_path = tmp_path / "base.json"
    base_path.write_text(json.dumps(base), encoding="utf-8")
    output = tmp_path / "FSH.profile.json"
    build_general_document_profile(
        project_id="FSH",
        workbook_path=workbook,
        base_profile_path=base_path,
        output_path=output,
    )
    profile = json.loads(output.read_text(encoding="utf-8"))
    assert profile["profile_name"] == "FSH"
    assert profile["workbook"]["expected_sheets"] == ["SUBSTATION"]
    assert profile["structure"]["m_line"]["target_sheets"] == ["SUBSTATION"]
    assert profile["known_exceptions"] == []
    assert profile["known_fixes"] == []
    assert "SUBSTATION" in profile["data_model"]["instances"]
