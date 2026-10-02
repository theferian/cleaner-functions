from pathlib import Path

from openpyxl import Workbook

from cleaner_core.overlay import ResolvedWorkbookView, apply_state_to_profile
from cleaner_core.snapshot import WorkbookSnapshot
from cleaner_core.state import new_resolved_state, ingest_review_decisions


def decisions_payload():
    return {
        "schema": "cleaner.review_decisions/0.1",
        "run": {"script_id": "M_LINE"},
        "decisions": [
            {
                "review_id": "r1",
                "source": {"test_id": "M_LINE"},
                "location": {"sheet": "S", "target": "A1:B1"},
                "decision": {
                    "action": "UNMERGE_RANGE",
                    "parameters": {"sheet": "S", "target": "A1:B1"},
                    "reason": "repair",
                    "decided_at": "2026-10-02T20:00:00",
                },
            },
            {
                "review_id": "r2",
                "source": {"test_id": "M_LINE"},
                "location": {"sheet": "S", "target": "C3"},
                "decision": {
                    "action": "SET_VALUE",
                    "parameters": {"sheet": "S", "target": "C3", "value": "1.2.1000"},
                    "reason": "logical repair",
                    "decided_at": "2026-10-02T20:00:01",
                },
            },
            {
                "review_id": "r3",
                "source": {"test_id": "M_LINE"},
                "location": {"sheet": "S", "target": "A10:B11"},
                "decision": {
                    "action": "ACCEPT_EXCEPTION",
                    "parameters": {
                        "test_id": "M_LINE",
                        "sheet": "S",
                        "target": "A10:B11",
                        "rows": [10, 11],
                    },
                    "reason": "accepted",
                    "decided_at": "2026-10-02T20:00:02",
                },
            },
        ],
    }


def test_state_ingests_reviewed_resolutions(tmp_path):
    state = new_resolved_state(run_id="run1", profile_name="AZD")
    ingest_review_decisions(state, decisions_payload())

    assert state["stages"]["M_LINE"]["status"] == "REVIEWED"
    assert state["value_overrides"]["S!C3"]["value"] == "1.2.1000"
    assert state["merge_operations"][0]["action"] == "UNMERGE_RANGE"
    assert state["accepted_exceptions"][0]["rows"] == [10, 11]


def test_overlay_changes_logical_workbook_and_profile(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "header"
    ws.merge_cells("A1:B1")
    ws["C3"] = "-"
    wb.save(path)

    base = WorkbookSnapshot(path)
    state = new_resolved_state(run_id="run1", profile_name="AZD")
    ingest_review_decisions(state, decisions_payload())
    view = ResolvedWorkbookView(base, state)

    assert view.merged_range_containing("S", "A1") is None
    assert view.value("S", "C3") == "1.2.1000"

    profile = apply_state_to_profile({"known_exceptions": []}, state)
    assert profile["known_exceptions"][0]["match"]["rows"] == [10, 11]
    base.close()
