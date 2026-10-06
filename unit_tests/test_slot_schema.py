from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from cleaner_core.card_map import scan_card_map
from cleaner_core.field_schema import (
    infer_slot_schema,
    infer_reference_regex,
    _detect_block_width,
    is_valid_jalali,
    normalize_jalali,
    normalize_embedded_jalali_text,
)
from cleaner_core.snapshot import WorkbookSnapshot


def make_profile():
    return {
        "profile_name": "AZD",
        "structure": {"m_line": {"target_sheets": ["230KV SS"], "boundary_after_col": 3}},
        "data_model": {
            "family_fields": [
                {"key": "document_type", "default_locator": {"kind": "COLUMN", "column": 1}},
                {"key": "name", "default_locator": {"kind": "COLUMN", "column": 2}},
                {"key": "wbs", "default_locator": {"kind": "COLUMN", "column": 3}},
            ],
            "instances": {"230KV SS": {"field_map": {
                "document_type": {"kind": "COLUMN", "column": 1},
                "name": {"kind": "COLUMN", "column": 2},
                "wbs": {"kind": "COLUMN", "column": 3},
            }}},
            "document_boundaries": {
                "first_data_row": 3,
                "signal_fields": ["document_type", "name", "wbs"],
                "boundary_fields": ["name", "wbs"],
                "separator_fields": ["name", "wbs"],
            },
            "card_map": {"data_start_col": 4},
            "field_inference": {
                "header_rows": [1, 2],
                "max_block_width": 6,
                "min_block_confidence": 0.65,
                "reference_min_samples": 3,
                "reference_min_coverage": 0.60,
                "status_min_occurrences": 2,
                "status_min_coverage": 0.70,
            },
        },
    }


def build_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for col, header in [(4, "Send Date"), (5, "Receive Date"), (6, "Status"),
                        (7, "Send Date"), (8, "Receive Date"), (9, "Status")]:
        ws.cell(2, col).value = header
    rows = [3, 5, 7, 9]
    statuses = ["CM", "RJ", "CM", "RJ"]
    for i, row in enumerate(rows):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
        for start_col, month in [(4, 1), (7, 2)]:
            ws.cell(row, start_col).value = f"1404/{month:02d}/{i+1:02d}"
            ws.cell(row, start_col + 1).value = f"1404/{month:02d}/{i+2:02d}"
            ws.cell(row, start_col + 2).value = statuses[i]
            ws.merge_cells(start_row=row, start_column=start_col + 2, end_row=row + 1, end_column=start_col + 2)
    wb.save(path)


def test_slot_schema_discovers_three_column_transaction_pattern(tmp_path):
    path = tmp_path / "book.xlsx"
    build_book(path)
    snap = WorkbookSnapshot(path)
    scan = scan_card_map(snap, make_profile(), "230KV SS")
    schema = infer_slot_schema(snap, make_profile(), "230KV SS", scan)
    assert schema.block_width == 3
    assert schema.block_confidence >= 0.65
    assert schema.roles_by_offset[0] == "send_date"
    assert schema.roles_by_offset[1] == "receive_date"
    assert schema.roles_by_offset[2] == "status"
    assert schema.status_values == ("CM", "RJ")
    assert schema.issues == ()
    snap.close()


def test_jalali_normalization_and_calendar_validation():
    assert normalize_jalali("۱۴۰۴-۲-۳") == "1404/02/03"
    assert normalize_jalali("04/06/03") == "1404/06/03"
    assert is_valid_jalali("1404/02/30") is True
    assert is_valid_jalali("1404/13/01") is False
    assert is_valid_jalali("1403/12/30") is True
    assert is_valid_jalali("1404/12/30") is False


def test_reference_regex_is_inferred_from_majority_shape():
    rule = infer_reference_regex(["AZD-ENG-1001", "AZD-ENG-1002", "AZD-ENG-1003", "BAD VALUE"], min_samples=3)
    assert rule is not None
    assert rule.coverage == 0.75
    assert rule.regex == r"^[A-Za-z]{3}\-[A-Za-z]{3}\-\d{4}$"


def test_block_width_prefers_fundamental_period_over_harmonic():
    # Width 6 is a perfect harmonic, but width 3 is almost as good and is the
    # fundamental transaction rhythm, so 3 must win.
    kinds = [
        "DATE", "DATE", "STATUS",
        "DATE", "MIXED", "STATUS",
        "DATE", "DATE", "STATUS",
        "DATE", "MIXED", "STATUS",
        "DATE", "DATE", "STATUS",
        "DATE", "MIXED", "STATUS",
    ]
    width, confidence = _detect_block_width(
        kinds, 8, fundamental_tolerance=0.20
    )
    assert width == 3
    assert confidence >= 0.70


def test_mixed_column_does_not_inherit_date_role(tmp_path):
    path = tmp_path / "mixed.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for col, header in [(4, "Send Date"), (5, "Receive Date"), (6, "Status"),
                        (7, "Send Date"), (8, "Receive Date"), (9, "Status")]:
        ws.cell(2, col).value = header

    rows = [3, 5, 7, 9, 11, 13]
    for i, row in enumerate(rows):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
        ws[f"D{row}"] = f"1404/01/{i+1:02d}"
        ws[f"E{row}"] = f"1404/01/{i+2:02d}"
        ws[f"F{row}"] = "CM"
        ws[f"G{row}"] = f"1404/02/{i+1:02d}"
        ws[f"H{row}"] = (
            f"1404/02/{i+2:02d}" if i < 3 else f"عطف به نامه شماره {100+i}"
        )
        ws[f"I{row}"] = "RJ"
        ws.merge_cells(start_row=row, start_column=6, end_row=row + 1, end_column=6)
        ws.merge_cells(start_row=row, start_column=9, end_row=row + 1, end_column=9)
    wb.save(path)

    profile = make_profile()
    profile["data_model"]["field_inference"]["semantic_min_confidence"] = 0.90
    profile["data_model"]["field_inference"]["block_fundamental_tolerance"] = 0.20
    snap = WorkbookSnapshot(path)
    scan = scan_card_map(snap, profile, "230KV SS")
    schema = infer_slot_schema(snap, profile, "230KV SS", scan)

    by_letter = {p.letter: p for p in schema.columns}
    assert by_letter["E"].semantic_type == "DATE"
    assert by_letter["E"].inferred_role == "receive_date"
    assert by_letter["H"].semantic_type == "MIXED"
    assert by_letter["H"].inferred_role is None
    assert schema.role_for_column(8) is None
    snap.close()


def test_vertical_lanes_separate_dates_from_references_in_same_column(tmp_path):
    path = tmp_path / "lane_schema.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for col, header in [(4, "Send Date"), (5, "Receive Date"), (6, "Status"),
                        (7, "Send Date"), (8, "Receive Date"), (9, "Status")]:
        ws.cell(2, col).value = header

    for i, row in enumerate([3, 5, 7, 9, 11, 13]):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
        # Same absolute columns carry two vertical semantic lanes.
        ws[f"D{row}"] = f"1404/01/{i+1:02d}"
        ws[f"D{row+1}"] = f"AZD-SND-{1000+i}"
        ws[f"E{row}"] = f"1404/01/{i+2:02d}"
        ws[f"E{row+1}"] = f"AZD-RCV-{2000+i}"
        ws[f"F{row}"] = "CM"
        ws.merge_cells(start_row=row, start_column=6, end_row=row + 1, end_column=6)

        ws[f"G{row}"] = f"1404/02/{i+1:02d}"
        ws[f"G{row+1}"] = f"AZD-SND-{3000+i}"
        ws[f"H{row}"] = f"1404/02/{i+2:02d}"
        ws[f"H{row+1}"] = f"AZD-RCV-{4000+i}"
        ws[f"I{row}"] = "RJ"
        ws.merge_cells(start_row=row, start_column=9, end_row=row + 1, end_column=9)
    wb.save(path)

    profile = make_profile()
    profile["data_model"]["field_inference"]["semantic_min_confidence"] = 0.90
    snap = WorkbookSnapshot(path)
    scan = scan_card_map(snap, profile, "230KV SS")
    schema = infer_slot_schema(snap, profile, "230KV SS", scan)

    slot_map = {(p.letter, p.lane): p for p in schema.slot_profiles}
    assert slot_map[("D", "TOP")].semantic_type == "DATE"
    assert slot_map[("D", "TOP")].inferred_role == "send_date"
    assert slot_map[("D", "BOTTOM")].semantic_type == "REFERENCE"
    assert slot_map[("D", "BOTTOM")].inferred_role == "reference"
    assert slot_map[("F", "FULL")].semantic_type == "STATUS"
    assert slot_map[("F", "FULL")].inferred_role == "status"
    assert schema.block_width == 3
    snap.close()


def test_four_row_card_normalizes_to_same_top_bottom_lanes(tmp_path):
    path = tmp_path / "four_row_lane_schema.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "230KV SS"
    for col, header in [(4, "Send Date"), (5, "Receive Date"), (6, "Status"),
                        (7, "Send Date"), (8, "Receive Date"), (9, "Status")]:
        ws.cell(2, col).value = header

    for i, row in enumerate([3, 7, 11, 15]):
        ws[f"A{row}"] = "DOC"
        ws[f"B{row}"] = f"Doc {i}"
        ws[f"C{row}"] = f"1.{i+1}"
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 3, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row + 3, end_column=3)
        for start_col, prefix in [(4, "SND"), (7, "RCV")]:
            ws.cell(row, start_col).value = f"1404/01/{i+1:02d}"
            ws.merge_cells(start_row=row, start_column=start_col, end_row=row + 1, end_column=start_col)
            ws.cell(row + 2, start_col).value = f"AZD-{prefix}-{1000+i}"
            ws.merge_cells(start_row=row + 2, start_column=start_col, end_row=row + 3, end_column=start_col)
            ws.cell(row, start_col + 1).value = f"1404/02/{i+1:02d}"
            ws.merge_cells(start_row=row, start_column=start_col + 1, end_row=row + 1, end_column=start_col + 1)
            ws.cell(row + 2, start_col + 1).value = f"AZD-{prefix}-{2000+i}"
            ws.merge_cells(start_row=row + 2, start_column=start_col + 1, end_row=row + 3, end_column=start_col + 1)
            ws.cell(row, start_col + 2).value = "CM"
            ws.merge_cells(start_row=row, start_column=start_col + 2, end_row=row + 3, end_column=start_col + 2)
    wb.save(path)

    profile = make_profile()
    snap = WorkbookSnapshot(path)
    scan = scan_card_map(snap, profile, "230KV SS")
    schema = infer_slot_schema(snap, profile, "230KV SS", scan)
    slot_map = {(p.letter, p.lane): p for p in schema.slot_profiles}
    assert slot_map[("D", "TOP")].inferred_role == "send_date"
    assert slot_map[("D", "BOTTOM")].inferred_role == "reference"
    assert slot_map[("F", "FULL")].inferred_role == "status"
    snap.close()


def test_embedded_short_jalali_is_normalized_inside_explanation():
    text, dates, invalid = normalize_embedded_jalali_text(
        "عطف به بند شماره 1 صورتجلسه مورخ 04/06/03"
    )
    assert text == "عطف به بند شماره 1 صورتجلسه مورخ 1404/06/03"
    assert dates == ("1404/06/03",)
    assert invalid == ()


def test_anchor_spacing_recovers_block_width_from_repeating_status_columns():
    from cleaner_core.field_schema import SlotProfile, _infer_block_width_from_semantic_anchors

    profiles = []
    # Real AZD 63KV pattern: F, I, L, O, R are stable status columns.
    for col, letter in [(6, "F"), (9, "I"), (12, "L"), (15, "O"), (18, "R")]:
        profiles.append(SlotProfile(
            column=col,
            letter=letter,
            lane="FULL",
            headers=(),
            non_empty=20,
            dominant_kind="CODE",
            dominant_kind_confidence=1.0,
            unique_count=5,
            samples=("CM", "RJ"),
            semantic_type="STATUS",
            semantic_confidence=1.0,
            inferred_role="status",
            role_confidence=1.0,
            role_source="MAJORITY_PATTERN",
        ))
    # Add noisy stable fields between anchors. These are exactly the kind of
    # evidence that can make whole-column periodicity weak even though the
    # transaction rhythm is obvious from status columns.
    for col in [4, 5, 7, 8, 10, 11, 13, 14, 16, 17]:
        profiles.append(SlotProfile(
            column=col,
            letter=get_column_letter(col),
            lane="TOP",
            headers=(),
            non_empty=10,
            dominant_kind="TEXT",
            dominant_kind_confidence=0.95,
            unique_count=10,
            samples=("note",),
            semantic_type="TEXT",
            semantic_confidence=0.95,
        ))

    width, confidence, evidence = _infer_block_width_from_semantic_anchors(
        profiles,
        max_width=8,
        min_semantic_confidence=0.90,
        min_columns=3,
        min_spacing_coverage=0.75,
    )
    assert width == 3
    assert confidence == 1.0
    assert evidence["semantic_anchor"] == "STATUS"
    assert evidence["columns"] == [6, 9, 12, 15, 18]


def test_mixed_offsets_gain_send_receive_roles_from_value_level_dates(tmp_path):
    from cleaner_core.field_schema import _augment_mixed_offset_date_roles
    from cleaner_core.card_map import CardMap, CardMapScan, CardSlot

    cards = []
    for idx in range(4):
        row = 3 + idx * 4
        slots = []
        for block_start in (4, 7, 10):
            # Same horizontal column stacks date/reference/date/reference values.
            slots.extend([
                CardSlot(f"{get_column_letter(block_start)}{row}", 0, 0, block_start, block_start, False, f"1404/01/0{idx+1}"),
                CardSlot(f"{get_column_letter(block_start)}{row+1}", 1, 1, block_start, block_start, False, f"1404/ES/S/{1000+idx}"),
                CardSlot(f"{get_column_letter(block_start+1)}{row}", 0, 0, block_start+1, block_start+1, False, f"1404/02/0{idx+1}"),
                CardSlot(f"{get_column_letter(block_start+1)}{row+1}", 1, 1, block_start+1, block_start+1, False, f"1404/ES/R/{2000+idx}"),
                CardSlot(f"{get_column_letter(block_start+2)}{row}:{get_column_letter(block_start+2)}{row+3}", 0, 3, block_start+2, block_start+2, True, "CM"),
            ])
        cards.append(CardMap(
            sheet="63KV SS", record_index=idx+1, anchor_row=row,
            start_row=row, end_row=row+3, identifier_range=f"A{row}:C{row+3}",
            data_range=f"D{row}:L{row+3}", data_start_col=4, data_end_col=12,
            slots=tuple(slots),
        ))
    scan = CardMapScan(sheet="63KV SS", cards=cards, data_start_col=4, data_end_col=12)
    roles_by_slot = {"2:FULL": "status"}
    conf_by_slot = {"2:FULL": 1.0}
    roles_by_offset = {2: "status"}
    conf_by_offset = {2: 1.0}

    _augment_mixed_offset_date_roles(
        scan,
        data_start_col=4,
        block_width=3,
        roles_by_slot=roles_by_slot,
        role_conf_by_slot=conf_by_slot,
        roles_by_offset=roles_by_offset,
        role_conf_by_offset=conf_by_offset,
        min_samples=3,
        min_share=0.20,
    )
    assert roles_by_offset[0] == "send_date"
    assert roles_by_offset[1] == "receive_date"
    assert roles_by_offset[2] == "status"
    assert roles_by_slot["0:VALUE"] == "send_date"
    assert roles_by_slot["1:VALUE"] == "receive_date"
